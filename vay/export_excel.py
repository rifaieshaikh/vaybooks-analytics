"""Write one .xlsx with 31-char sheet names and merged FY headers. Not DataFrame.to_excel for wide sheets."""

from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from vay.config import EXCEL_SHEET_NAMES, NUMERIC_FORMAT, STATUS_COLORS, TEXT_HEADERS
from vay.dates import number_, round2

ILLEGAL = re.compile(r"[\\/?*\[\]:]")


def excel_sheet_name(report_id, used):
    if report_id in EXCEL_SHEET_NAMES:
        base = EXCEL_SHEET_NAMES[report_id]
    elif report_id.startswith("sales_follow_up_"):
        group = report_id[len("sales_follow_up_"):]
        base = _followup_name("sf_", group)
    elif report_id.startswith("collection_follow_up_"):
        group = report_id[len("collection_follow_up_"):]
        base = _followup_name("cf_", group)
    else:
        base = ILLEGAL.sub("-", report_id)[:31]
    base = (base or "report")[:31]
    name = base
    n = 2
    while name.lower() in used:
        suffix = "_%d" % n
        name = (base[: 31 - len(suffix)] + suffix)[:31]
        n += 1
    used.add(name.lower())
    return name


def sheet_names_by_report_id(reports):
    """Map report id → Excel sheet title using the same order as write_workbook."""
    used = set()
    out = {}
    for report in reports or []:
        rid = (report or {}).get("id") or ""
        if not rid:
            continue
        out[rid] = excel_sheet_name(rid, used)
    return out


def workbook_bytes_keep_ids(xlsx_bytes, all_reports, keep_ids):
    """Return a copy of the workbook with only sheets for keep_ids. None if empty."""
    from openpyxl import load_workbook

    keep_ids = set(keep_ids or [])
    if not keep_ids:
        return None
    names = sheet_names_by_report_id(all_reports)
    keep_names = {names[rid] for rid in keep_ids if rid in names}
    if not keep_names:
        return None
    wb = load_workbook(BytesIO(xlsx_bytes))
    for ws in list(wb.worksheets):
        if ws.title not in keep_names:
            wb.remove(ws)
    if not wb.worksheets:
        return None
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _followup_name(prefix, group):
    """Truncate so _2 still fits in 31 characters."""
    cleaned = ILLEGAL.sub("-", str(group))
    max_group = 31 - len(prefix) - 2
    return (prefix + cleaned[:max_group])[:31]


def _is_numeric_header(header):
    return header not in TEXT_HEADERS and header not in ("", None)


def write_workbook(reports):
    wb = Workbook()
    default = wb.active
    used = set()
    first = True
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for report in reports:
        name = excel_sheet_name(report["id"], used)
        ws = default if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        wide = report.get("header_row1")
        start_row = 1
        if wide:
            for c, val in enumerate(report["header_row1"], 1):
                cell = ws.cell(1, c, val)
                cell.font = Font(bold=True)
            for c, val in enumerate(report["headers"], 1):
                cell = ws.cell(2, c, val)
                cell.font = Font(bold=True)
            for a, b, _text in report.get("merges") or []:
                if b > a:
                    ws.merge_cells(start_row=1, start_column=a + 1, end_row=1, end_column=b + 1)
                    ws.cell(1, a + 1).alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A3"
            start_row = 3
        else:
            for c, val in enumerate(report["headers"], 1):
                cell = ws.cell(1, c, val)
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            start_row = 2
        status_col = None
        if report.get("status_header") and report["status_header"] in report["headers"]:
            status_col = report["headers"].index(report["status_header"]) + 1
        r = start_row
        body = list(report.get("rows") or [])
        if report.get("total"):
            body.append(report["total"])
        headers = report["headers"]
        for row in body:
            for c, val in enumerate(row, 1):
                header = headers[c - 1] if c - 1 < len(headers) else ""
                cell = ws.cell(r, c, val if val != "" else None)
                cell.border = thin
                if val != "" and val is not None and _is_numeric_header(header):
                    try:
                        num = number_(val) if not isinstance(val, (int, float)) else float(val)
                        cell.value = round2(num)
                        cell.number_format = NUMERIC_FORMAT
                    except Exception:
                        pass
                if status_col and c == status_col:
                    color = STATUS_COLORS.get(str(val), "")
                    if color:
                        cell.fill = PatternFill("solid", fgColor=color.replace("#", ""))
            r += 1
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(max(1, len(headers))), max(1, r - 1))
        ws.sheet_view.showGridLines = False
    if first:
        default.title = "empty"
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
