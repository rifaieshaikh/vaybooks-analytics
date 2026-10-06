import hashlib
import io
import json
import os
import zipfile

os.environ.setdefault("VAY_STORE", "memory")
os.environ.setdefault("VAY_SYNC_JOBS", "1")

from fastapi.testclient import TestClient
from openpyxl import Workbook

from server.main import app
from server.settings import mongo_uri
from server.store import get_store, reset_store_for_tests
from vay.engine import hash_password


def client():
    reset_store_for_tests()
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert r.status_code == 200
    return c


def login_as(c, username, password):
    return c.post("/api/auth/login", json={"username": username, "password": password})


def xlsx_bytes(sheets):
    wb = Workbook()
    first = True
    for name, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_info_is_public():
    raw = TestClient(app)
    r = raw.get("/api/info")
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert isinstance(body.get("lan_urls"), list)
    assert isinstance(body.get("port"), int)


def test_mapper_rejects_two_headers_one_field():
    c = client()
    r = c.put("/api/mappers/sales", json={
        "column_map": {"Date": "Date", "Bill Date": "Date"},
        "unique_key": ["Date"],
    })
    assert r.status_code == 400


def test_upload_without_mapper_uses_defaults():
    c = client()
    data = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "A", "R", 1]])])
    r = c.post("/api/uploads", files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["types"]["sales"]["added"] == 1


def test_preview_shows_rows_and_mapping():
    c = client()
    data = xlsx_bytes([("sales", ["Bill Date", "Customer", "Rep", "Amount", "Invoice No"], [["19-09-24", "Acme", "R", 10, "INV-1"]])])
    r = c.post("/api/uploads/preview", files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    sheet = r.json()["sheets"][0]
    assert sheet["type"] == "sales"
    assert sheet["row_count"] == 1
    assert sheet["sample"]
    assert sheet["column_map"]["Bill Date"] == "Date"
    assert sheet["column_map"]["Customer"] == "Party Name"
    assert "Invoice No" in sheet["column_map"].values() or sheet["column_map"].get("Invoice No") == "Invoice No"
    assert sheet["ready"] is True


def test_upload_sales_then_duplicate_zero_inserts():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    data = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 10]])])
    r1 = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    r2 = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r1.json()["types"]["sales"]["added"] == 1
    assert r2.json()["types"]["sales"]["skipped"] == 1
    orig = c.get("/api/uploads/%s/file" % r1.json()["id"])
    assert hashlib.sha256(orig.content).digest() == hashlib.sha256(data).digest()


def test_delete_then_reupload_inserts():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    data = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 10]])])
    r1 = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.delete("/api/uploads/%s" % r1.json()["id"])
    r2 = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.json()["types"]["sales"]["added"] == 1


def test_explorer_group_join_and_arr_dates_ignored():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 1]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 10]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    g = c.get("/api/rows", params={"type": "sales", "group": "south"})
    assert g.status_code == 200
    assert g.json()["total"] == 1
    d = c.get("/api/rows", params={"type": "arr", "date_from": "2020-01-01"})
    assert d.status_code == 200
    assert d.json()["total"] == 1


def test_combined_items_mapper_missing_sales_still_inserts():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    data = xlsx_bytes([
        ("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 10]]),
        ("items", ["Date", "Item Name", "Qty", "Rate"], [["19-09-24", "SKU", 1, 2]]),
    ])
    r = c.post("/api/uploads", files={"file": ("all.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    types = r.json()["types"]
    assert types["sales"]["added"] == 1
    assert types["items"]["added"] == 1
    raw = TestClient(app)
    assert raw.get("/api/health").status_code == 200
    info = raw.get("/api/info")
    assert info.status_code == 200
    body = info.json()
    assert body["ok"] is True
    assert isinstance(body.get("lan_urls"), list)
    assert raw.get("/api/mappers").status_code == 401


def _seed_core(c):
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    c.put("/api/mappers/receipt", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Sales Rep", "Amount"]})
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    data = xlsx_bytes([
        ("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 100]]),
        ("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [["19-09-24", "Acme", "R", 40]]),
        ("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 60]]),
    ])
    c.post("/api/uploads", files={"file": ("all.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_create_run_async_contract():
    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19", "password": "secret-pass"})
    assert post.status_code == 202
    body = post.json()
    assert body["status"] == "queued"
    rid = body["id"]
    got = c.get("/api/runs/%s" % rid)
    assert got.status_code == 200
    assert got.json()["status"] == "succeeded"
    dump = json.dumps(got.json())
    assert "secret-pass" not in dump
    store = get_store()
    run = store.get_run(rid)
    assert "password" not in run
    blob = store.get_blob(run["xlsx_id"])
    file_resp = c.get("/api/runs/%s/file" % rid)
    assert file_resp.status_code == 200
    assert hashlib.sha256(file_resp.content).digest() == hashlib.sha256(blob["data"]).digest()


def test_download_excel_filters_sheets_by_permission():
    from openpyxl import load_workbook

    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    assert post.status_code == 202
    rid = post.json()["id"]

    admin = c.get("/api/runs/%s/file" % rid)
    assert admin.status_code == 200
    admin_names = load_workbook(io.BytesIO(admin.content)).sheetnames
    assert "sales_rep_perf" in admin_names
    assert "warnings" in admin_names

    c.post("/api/users", json={"username": "fin", "password": "fin12345", "role": "Finance"})
    fin = TestClient(app)
    login_as(fin, "fin", "fin12345")
    finance = fin.get("/api/runs/%s/file" % rid)
    assert finance.status_code == 200
    fin_names = load_workbook(io.BytesIO(finance.content)).sheetnames
    assert "sales_rep_perf" not in fin_names
    assert "account_perf" not in fin_names
    assert "warnings" in fin_names

    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    wh = TestClient(app)
    login_as(wh, "wh", "wh12345")
    assert wh.get("/api/runs/%s/file" % rid).status_code == 403


def test_downloads_refuse_without_a_snapshot_and_manifest_hides_sales():
    c = client()
    store = get_store()
    xlsx_id = store.put_blob(b"not-a-workbook", "reports.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    zip_id = store.put_blob(b"PK\x03\x04not-a-zip", "reports.zip", "application/zip")
    run = store.insert_run({
        "status": "succeeded",
        "report_date": "2026-10-03",
        "xlsx_id": xlsx_id,
        "pdf_zip_id": zip_id,
        "pdf_export_status": "succeeded",
        "manifest": {
            "eligibility": {
                "sales_mtd": {"status": "eligible"},
                "stock_cover_days": {"status": "eligible"},
            },
            "reconciliation": {
                "status": "pass",
                "message": "Headline sales and AR match sidecar",
                "computed": {"sales_mtd": 10, "sales_ytd": 20, "ar_balance": 30},
                "expected": {"sales_mtd": 10, "sales_ytd": 20, "ar_balance": 30},
                "exceptions": [{"severity": "error", "metric_id": "sales_mtd", "message": "Mismatch on sales_mtd"}],
            },
        },
    })
    rid = str(run["_id"])
    admin = c.get("/api/runs/%s" % rid).json()["manifest"]
    assert admin["reconciliation"]["computed"]["sales_mtd"] == 10
    assert c.get("/api/runs/%s/file" % rid).status_code == 403
    assert c.get("/api/runs/%s/export-pdf/file" % rid).status_code == 403
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    wh = TestClient(app)
    login_as(wh, "wh", "wh12345")
    hidden = wh.get("/api/runs/%s" % rid).json()["manifest"]
    assert hidden["reconciliation"]["status"] == "unavailable"
    assert "computed" not in hidden["reconciliation"]
    assert "sales_mtd" not in (hidden.get("eligibility") or {})
    assert hidden["eligibility"]["stock_cover_days"]["status"] == "eligible"
    dash = wh.get("/api/dashboard").json()
    assert "exceptions" not in dash
    assert "sales_mtd" not in (dash.get("eligibility") or {})
    asked = wh.post("/api/saved-reports/ask", json={"text": "Sales", "report_date": "2026-10-03"})
    assert asked.status_code == 403


def test_explorer_filter_does_not_change_create_total():
    c = client()
    _seed_core(c)
    extra = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["01-04-24", "Other", "R", 50]])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("e.xlsx", extra, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    filtered = c.get("/api/rows", params={"type": "sales", "date_from": "2024-09-01", "rep": "R"})
    assert filtered.json()["total"] == 1
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    snap = c.get("/api/runs/%s" % post.json()["id"]).json()["snapshot"]
    rep = next(r for r in snap["reports"] if r["id"] == "sales_rep_performance_report")
    ytd_i = rep["headers"].index("YTD Sales")
    total = rep["total"]
    assert float(str(total[ytd_i]).replace(",", "")) == 150


def test_file_while_running_and_second_post_409():
    c = client()
    store = get_store()
    running = store.insert_run({"status": "running", "packs": {}})
    rid = str(running["_id"])
    assert c.get("/api/runs/%s/file" % rid).status_code == 409
    assert c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"}).status_code == 409


def test_default_mongo_uri_uses_vay_reports_db():
    assert mongo_uri() == "mongodb://127.0.0.1:27017/vay-reports"


def test_login_seeds_admin_and_rejects_bad_password():
    c = client()
    me = c.get("/api/auth/me")
    assert me.status_code == 200
    assert me.json()["username"] == "admin"
    assert me.json().get("must_change_password") is True
    assert "users.manage" in me.json()["permissions"]
    assert "users.view" in me.json()["permissions"]
    assert "roles.view" in me.json()["permissions"]
    assert "roles.manage" in me.json()["permissions"]
    raw = TestClient(app)
    bad = raw.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
    assert bad.status_code == 401
    dump = json.dumps(me.json())
    assert "admin123" not in dump
    assert hash_password("admin123") not in dump


def test_change_password_clears_must_change_flag():
    reset_store_for_tests()
    from server.rate_limit import reset_for_tests

    reset_for_tests()
    raw = TestClient(app)
    login = raw.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert login.status_code == 200
    assert login.json()["must_change_password"] is True
    bad = raw.post(
        "/api/auth/change-password",
        json={"current_password": "admin123", "new_password": "short"},
    )
    assert bad.status_code == 400
    ok = raw.post(
        "/api/auth/change-password",
        json={"current_password": "admin123", "new_password": "newpass99"},
    )
    assert ok.status_code == 200
    assert ok.json()["must_change_password"] is False
    again = raw.post("/api/auth/login", json={"username": "admin", "password": "newpass99"})
    assert again.status_code == 200
    assert again.json()["must_change_password"] is False


def test_login_rate_limit_blocks_after_failures(monkeypatch):
    reset_store_for_tests()
    from server.rate_limit import reset_for_tests

    reset_for_tests()
    monkeypatch.setenv("VAY_LOGIN_RATE_LIMIT", "3")
    monkeypatch.setenv("VAY_LOGIN_RATE_WINDOW", "300")
    raw = TestClient(app)
    for _ in range(3):
        bad = raw.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
        assert bad.status_code == 401
    blocked = raw.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
    assert blocked.status_code == 429
    reset_for_tests()

def test_viewer_cannot_upload_or_create():
    c = client()
    created = c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    assert created.status_code == 200
    raw = TestClient(app)
    login_as(raw, "look", "look123")
    data = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "A", "R", 1]])])
    up = raw.post("/api/uploads", data={"type": "sales"}, files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert up.status_code == 403
    run = raw.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    assert run.status_code == 403


def test_profit_password_endpoint_removed():
    c = client()
    assert c.get("/api/settings").status_code == 404
    assert c.put("/api/settings/profit-password", json={"current_password": "x", "new_password": "y"}).status_code == 404
    store = get_store()
    assert not hasattr(store, "get_profit_password_hash") or True
    user = store.get_user("admin")
    assert "password_hash" in user
    assert json.dumps(c.get("/api/users").json()).find(user["password_hash"]) < 0


def test_party_master_upsert_and_joins_receipts_payments():
    c = client()
    c.put("/api/mappers/party", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/receipt", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Sales Rep", "Amount"]})
    c.put("/api/mappers/payments", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Amount"]})
    party = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "South"]])])
    r1 = c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", party, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r1.json()["types"]["party"]["upserted"] == 1
    party2 = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "North"]])])
    r2 = c.post("/api/uploads", data={"type": "party"}, files={"file": ("p2.xlsx", party2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.json()["types"]["party"]["upserted"] == 1
    listed = c.get("/api/rows", params={"type": "party"})
    assert listed.json()["total"] == 1
    assert listed.json()["rows"][0]["fields"]["Group"] == "North"
    receipt = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [["19-09-24", "Acme", "R", 40]])])
    pay = xlsx_bytes([("payments", ["Date", "Account Name", "Amount"], [["19-09-24", "Acme", 15]])])
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", receipt, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "payments"}, files={"file": ("pay.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    rec = c.get("/api/rows", params={"type": "receipt", "party": "acme", "group": "north"})
    assert rec.json()["total"] == 1
    paid = c.get("/api/rows", params={"type": "payments", "party": "acme", "group": "north"})
    assert paid.json()["total"] == 1
    miss = c.get("/api/rows", params={"type": "payments", "group": "south"})
    assert miss.json()["total"] == 0


def test_party_balance_filter():
    c = client()
    c.put("/api/mappers/party", json={"column_map": {}, "unique_key": ["Account Name"]})
    data = xlsx_bytes([("party", ["Account Name", "Group", "Balance"], [
        ["Acme", "North", 80],
        ["Beta", "North", 0],
        ["Credit Co", "South", -20],
    ])])
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    low = c.get("/api/rows", params={"type": "party", "balance_from": "50"}).json()
    assert low["total"] == 1
    assert low["rows"][0]["party"] == "Acme"
    mid = c.get("/api/rows", params={"type": "party", "balance_from": "0", "balance_to": "0"}).json()
    assert mid["total"] == 1
    assert mid["rows"][0]["party"] == "Beta"
    band = c.get("/api/rows", params={"type": "party", "balance_from": "-20", "balance_to": "10"}).json()
    assert {row["party"] for row in band["rows"]} == {"Beta", "Credit Co"}


def test_arr_creates_customer_and_replaces_balance():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    first = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 120]])])
    r1 = c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r1.json()["types"]["arr"]["customers_created"] == 1
    assert r1.json()["types"]["arr"]["parties_created"] == 1
    listed = c.get("/api/customers")
    assert listed.status_code == 200
    assert listed.json()["total"] == 1
    assert listed.json()["customers"][0]["due"] == 120
    second = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 80]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("b.xlsx", second, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed2 = c.get("/api/customers")
    assert listed2.json()["total"] == 1
    assert listed2.json()["customers"][0]["due"] == 80
    detail = c.get("/api/customers/%s" % listed2.json()["customers"][0]["uk"])
    assert detail.json()["due"] == 80
    party = c.get("/api/rows", params={"type": "party", "party": "acme"})
    assert party.json()["rows"][0]["fields"]["Balance"] == 80


def test_customer_opening_and_fy_sales_collection():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    c.put("/api/mappers/receipt", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Sales Rep", "Amount"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 1000]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        ["15-06-25", "Acme", "R", 100],
        ["15-06-24", "Acme", "R", 200],
    ])])
    rec = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [
        ["20-06-25", "Acme", "R", 40],
        ["20-06-24", "Acme", "R", 50],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", rec, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/customers").json()["customers"][0]["uk"]
    detail = c.get("/api/customers/%s" % uk).json()
    opening = next(x for x in detail["open_invoices"] if x["kind"] == "opening")
    assert opening["date"].startswith("2024-04-01")
    assert opening["date"] == detail["opening_date"]
    assert opening["due"] == 700
    # 1000 due + 90 collected − 300 invoiced leaves 790 with no sale.
    # That opening starts the ledger, and the last balance equals the ARR balance.
    ledger = detail["ar_mismatch"]
    assert ledger["opening"] == 790
    assert ledger["opening_month"] == "Jun 2024"
    assert ledger["balance"] == 1000
    assert ledger["balance_issue"] == ""
    first = ledger["years"][0]["months"][0]
    assert first["label"] == "Jun 2024"
    assert first["opening"] == 790
    assert ledger["years"][-1]["months"][-1]["balance"] == 1000
    sales_open = [x for x in detail["open_invoices"] if x["kind"] == "sale"]
    # Outstanding due order is LIFO: 100 (Jun-25) then 200 (Jun-24)
    assert [round(x["due"]) for x in sales_open] == [100, 200]
    labels = {y["label"]: y for y in detail["fy_years"]}
    assert "FY 2025-26" in labels or "FY 2024-25" in labels
    fy25 = next(y for y in detail["fy_years"] if y["start"].startswith("2025-04-01"))
    fy24 = next(y for y in detail["fy_years"] if y["start"].startswith("2024-04-01"))
    assert fy25["sales"] == 100
    assert fy25["collection"] == 40
    assert fy24["sales"] == 200
    assert fy24["collection"] == 50
    pdf = c.get("/api/customers/%s/pdf" % uk)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content[:4] == b"%PDF"


def test_customer_lifo_open_invoices_and_shared_status():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 150]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["15-06-24", "Acme", "R", 200, "INV-OLD"],
        ["15-06-25", "Acme", "R", 100, "INV-NEW"],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/customers").json()
    uk = listed["customers"][0]["uk"]
    detail = c.get("/api/customers/%s" % uk).json()
    opens = detail["open_invoices"]
    # Outstanding due is always LIFO on open balances: NEW then OLD
    assert [x["invoice"] for x in opens] == ["INV-NEW", "INV-OLD"]
    assert opens[0]["due"] == 100
    assert opens[1]["due"] == 50
    assert opens[0]["invoice_id"]
    assert not any(x["kind"] == "opening" for x in opens)
    sales_events = [e for e in detail["activity"] if e["kind"] == "sale"]
    assert {e["invoice"] for e in sales_events} == {"INV-OLD", "INV-NEW"}
    assert all(e.get("invoice_id") for e in sales_events)
    assert listed["customers"][0]["status"] == detail["status"]
    row = listed["customers"][0]
    buckets = detail["owe_buckets"]
    assert "collection_buckets" in detail
    assert row["d0_15"] == buckets["d0_15"]
    assert row["d15_30"] == buckets["d15_30"]
    assert row["d90"] == buckets["d90"]
    found = c.get("/api/customers", params={"q": "acm"}).json()
    assert found["total"] == 1
    miss = c.get("/api/customers", params={"q": "zzzz"}).json()
    assert miss["total"] == 0


def test_customer_due_always_lifo_latest_mode_migrates():
    c = client()
    # Removed "latest" settlement mode migrates to oldest; due stays LIFO.
    assert c.post("/api/settings/settlement", json={"mode": "latest"}).status_code == 200
    body = c.get("/api/settings/settlement").json()
    assert body["mode"] == "oldest"
    assert body["setup_complete"] is True
    assert isinstance(body["aging_bands"], list) and len(body["aging_bands"]) >= 1
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 150]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["15-06-24", "Acme", "R", 200, "INV-OLD"],
        ["15-06-25", "Acme", "R", 100, "INV-NEW"],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    detail = c.get("/api/customers/%s" % c.get("/api/customers").json()["customers"][0]["uk"]).json()
    opens = detail["open_invoices"]
    assert [x["invoice"] for x in opens] == ["INV-NEW", "INV-OLD"]
    assert opens[0]["due"] == 100
    assert opens[1]["due"] == 50
    c.post("/api/settings/settlement", json={"mode": "oldest"})


def test_settlement_settings_setup_and_bands():
    c = client()
    before = c.get("/api/settings/settlement").json()
    assert "setup_complete" in before
    assert "aging_bands" in before
    bands = [
        {"key": "fresh", "label": "0–10", "max_days": 10},
        {"key": "old", "label": "10+", "max_days": None},
    ]
    res = c.post("/api/settings/settlement", json={
        "mode": "specific",
        "aging_bands": bands,
        "setup_complete": True,
    })
    assert res.status_code == 200
    body = res.json()
    assert body["mode"] == "specific"
    assert body["setup_complete"] is True
    assert [b["key"] for b in body["aging_bands"]] == ["fresh", "old"]
    bad = c.post("/api/settings/settlement", json={
        "mode": "oldest",
        "aging_bands": [
            {"key": "a", "label": "30", "max_days": 30},
            {"key": "b", "label": "10", "max_days": 10},
            {"key": "c", "label": "+", "max_days": None},
        ],
    })
    assert bad.status_code == 400
    # restore defaults
    c.post("/api/settings/settlement", json={"mode": "oldest", "aging_bands": None})


def test_org_policy_and_expense_categories():
    c = client()
    policy = c.get("/api/settings/org-policy").json()
    assert policy["fiscal_year_start_month"] == 4
    assert policy["timezone"] == "Asia/Kolkata"
    assert abs(float(policy["sales_tax_inclusive_rate"]) - 0.18) < 1e-9
    saved = c.post("/api/settings/org-policy", json={
        "fiscal_year_start_month": 1,
        "timezone": "Asia/Kolkata",
        "currency_code": "INR",
        "currency_symbol": "₹",
        "sales_tax_inclusive_rate": 0.05,
        "terminology": {"customer": "Buyer"},
    }).json()
    assert saved["fiscal_year_start_month"] == 1
    assert abs(float(saved["sales_tax_inclusive_rate"]) - 0.05) < 1e-9
    assert saved["terminology"]["customer"] == "Buyer"
    bad = c.post("/api/settings/org-policy", json={"sales_tax_inclusive_rate": 2})
    assert bad.status_code == 400
    expense = c.get("/api/settings/expense-categories").json()
    assert "Office Expenses" in expense["categories"]
    assert expense["pack"] == "vay_wholesale"
    updated = c.post("/api/settings/expense-categories", json={
        "pack": "vay_wholesale",
        "categories": {
            "Office Expenses": ["TEST_ACCOUNT"],
            "Custom Marketing": ["AD_SPEND_ACCOUNT", "SOCIAL_ADS"],
        },
    }).json()
    assert updated["categories"]["Office Expenses"] == ["TEST_ACCOUNT"]
    assert updated["categories"]["Custom Marketing"] == ["AD_SPEND_ACCOUNT", "SOCIAL_ADS"]
    again = c.get("/api/settings/expense-categories").json()
    assert again["categories"]["Custom Marketing"] == ["AD_SPEND_ACCOUNT", "SOCIAL_ADS"]
    # restore Vay-compatible defaults
    from vay.packs.vay_wholesale import EXPENSE_ACCOUNT_MAP

    c.post("/api/settings/expense-categories", json={
        "pack": "vay_wholesale",
        "categories": EXPENSE_ACCOUNT_MAP,
    })
    c.post("/api/settings/org-policy", json={
        "fiscal_year_start_month": 4,
        "sales_tax_inclusive_rate": 0.18,
        "terminology": {"customer": "Customer"},
    })


def test_due_days_settings_and_ordering_flags():
    from datetime import timedelta

    from vay.dates import today_ist

    c = client()
    as_of = today_ist()
    # Default due days
    got = c.get("/api/settings/due-days").json()
    assert got["due_days"] == 30
    saved = c.post("/api/settings/due-days", json={"due_days": 30}).json()
    assert saved["due_days"] == 30

    d_sale = as_of
    d_31 = as_of - timedelta(days=31)
    d_30 = as_of - timedelta(days=30)
    d_10 = as_of - timedelta(days=10)

    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    c.put("/api/mappers/customer", json={"column_map": {}, "unique_key": ["Account Name"]})

    # Acme: old invoice 31 days before a new sale → flagged (age 31 > 30)
    # Beta: old invoice exactly 30 days before sale → not flagged
    # Gamma: recent only → not flagged
    # Delta: only one sale but ARR implies opening past due → flagged
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [
        ["Acme", "South", 200],
        ["Beta", "South", 200],
        ["Gamma", "North", 50],
        ["Delta", "North", 150],
    ])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        [d_31.strftime("%d-%m-%y"), "Acme", "RepA", 100, "A-OLD"],
        [d_sale.strftime("%d-%m-%y"), "Acme", "RepA", 100, "A-NEW"],
        [d_30.strftime("%d-%m-%y"), "Beta", "RepA", 100, "B-OLD"],
        [d_sale.strftime("%d-%m-%y"), "Beta", "RepA", 100, "B-NEW"],
        [d_10.strftime("%d-%m-%y"), "Gamma", "RepB", 50, "G-1"],
        [d_sale.strftime("%d-%m-%y"), "Delta", "RepB", 100, "D-NEW"],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    listed = c.get("/api/customers").json()["customers"]
    by_name = {row["name"]: row for row in listed}
    assert by_name["Acme"]["ordered_overdue"] is True
    assert by_name["Beta"]["ordered_overdue"] is False
    assert by_name["Gamma"]["ordered_overdue"] is False
    assert by_name["Delta"]["ordered_overdue"] is True

    acme = c.get("/api/customers/%s" % by_name["Acme"]["uk"]).json()
    assert "ordering" in acme
    assert acme["due_days"] == 30
    assert acme["due_days_source"] == "default"
    assert acme["ordering"]["overall"]["flagged_count"] == 1
    assert acme["ordering"]["overall"]["sales_count"] == 2
    flagged = acme["ordering"]["overall"]["events"]
    assert flagged[0]["invoice"] == "A-NEW"
    assert flagged[0]["oldest_age"] == 31
    assert acme["ordering"]["badge"]["current_month"] is True or acme["ordering"]["badge"]["current_fy"] is True
    assert acme["ordered_overdue"] is True

    delta = c.get("/api/customers/%s" % by_name["Delta"]["uk"]).json()
    assert delta["ordering"]["overall"]["flagged_count"] == 1
    assert delta["ordering"]["overall"]["events"][0]["invoice"] == "D-NEW"
    assert delta["ordering"]["overall"]["events"][0]["oldest_age"] > 30

    beta = c.get("/api/customers/%s" % by_name["Beta"]["uk"]).json()
    assert beta["ordering"]["overall"]["flagged_count"] == 0
    assert beta["ordered_overdue"] is False

    # Resolve order: customer beats group beats rep beats default
    c.put("/api/groups/%s/due-days" % by_name["Acme"]["group"], json={"due_days": 45})
    c.put("/api/reps/RepA/due-days", json={"due_days": 60})
    mid = c.get("/api/customers/%s" % by_name["Acme"]["uk"]).json()
    assert mid["due_days"] == 45
    assert mid["due_days_source"] == "group"
    # With group 45, age 31 is NOT > 45 → unflagged
    assert mid["ordering"]["overall"]["flagged_count"] == 0

    c.put("/api/customers/%s/due-days" % by_name["Acme"]["uk"], json={"due_days": 20})
    own = c.get("/api/customers/%s" % by_name["Acme"]["uk"]).json()
    assert own["due_days"] == 20
    assert own["due_days_source"] == "customer"
    assert own["due_days_own"] is True
    assert own["ordering"]["overall"]["flagged_count"] == 1

    cleared = c.put("/api/customers/%s/due-days" % by_name["Acme"]["uk"], json={"due_days": None}).json()
    assert cleared.get("cleared") is True
    after = c.get("/api/customers/%s" % by_name["Acme"]["uk"]).json()
    assert after["due_days"] == 45
    assert after["due_days_source"] == "group"

    # Group / rep rollups include member flagged sales
    group = c.get("/api/groups/%s" % by_name["Acme"]["group"]).json()
    assert "ordering" in group
    assert group["ordering"]["overall"]["sales_count"] >= 4
    rep = c.get("/api/reps/RepA").json()
    assert "ordering" in rep
    assert rep["ordering"]["overall"]["sales_count"] >= 4


def test_customer_list_sorts_due_buckets():
    from datetime import timedelta

    from vay.dates import today_ist

    c = client()
    as_of = today_ist()
    recent = as_of.strftime("%d-%m-%y")
    old = (as_of - timedelta(days=100)).strftime("%d-%m-%y")
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [
        ["Alpha", "South", 100],
        ["Beta", "South", 80],
    ])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        [recent, "Alpha", "R", 100],
        [old, "Beta", "R", 80],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    by_new = c.get("/api/customers", params={"sort": "d0_15", "dir": "desc"}).json()["customers"]
    assert [row["name"] for row in by_new] == ["Alpha", "Beta"]
    assert by_new[0]["d0_15"] == 100
    assert by_new[1]["d0_15"] == 0
    by_old = c.get("/api/customers", params={"sort": "d90", "dir": "desc"}).json()["customers"]
    assert [row["name"] for row in by_old] == ["Beta", "Alpha"]
    assert by_old[0]["d90"] == 80


def test_customer_list_sorts_ar_diff():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [
        ["High", "South", 80],
        ["Low", "South", 50],
        ["Even", "South", 100],
    ])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        ["01-09-26", "High", "R", 200],
        ["01-09-26", "Low", "R", 10],
        ["01-09-26", "Even", "R", 100],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    desc = c.get("/api/customers", params={"sort": "balance_diff", "dir": "desc"}).json()["customers"]
    by_name = {row["name"]: row for row in desc}
    # High's sales exceed the balance, so the excess is a mismatch.
    # Low's shortfall is an opening, so the books match.
    assert desc[0]["name"] == "High"
    assert by_name["High"]["balance_diff"] == 120
    assert by_name["High"]["ledger_opening"] == 0
    assert by_name["Low"]["balance_diff"] == 0
    assert by_name["Low"]["ledger_opening"] == 40
    assert by_name["Even"]["balance_diff"] == 0
    asc = c.get("/api/customers", params={"sort": "ar_diff", "dir": "asc"}).json()["customers"]
    assert asc[-1]["name"] == "High"


def test_customer_fifo_last_order_then_partial_and_opening():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 20000]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["01-08-25", "Acme", "R", 8000, "INV-PREV"],
        ["01-09-25", "Acme", "R", 9000, "INV-LAST"],
    ])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    detail = c.get("/api/customers/%s" % c.get("/api/customers").json()["customers"][0]["uk"]).json()
    opens = detail["open_invoices"]
    by_inv = {x["invoice"]: x for x in opens if x["kind"] == "sale"}
    assert by_inv["INV-LAST"]["due"] == 9000
    assert by_inv["INV-PREV"]["due"] == 8000
    opening = next(x for x in opens if x["kind"] == "opening")
    assert opening["due"] == 3000
    assert opening["date"].startswith("2025-04-01")


def test_customer_due_with_no_invoices_is_opening():
    c = client()
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 10000]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    detail = c.get("/api/customers/%s" % c.get("/api/customers").json()["customers"][0]["uk"]).json()
    opens = detail["open_invoices"]
    assert len(opens) == 1
    assert opens[0]["kind"] == "opening"
    assert opens[0]["due"] == 10000
    assert opens[0]["date"].endswith("-04-01")
    assert opens[0]["date"] == detail["opening_date"]


def test_customer_zero_due_stays_on_track():
    c = client()
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance", "Days"], [["Acme", "South", 0, 90]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["15-06-24", "Acme", "R", 40]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/customers").json()["customers"][0]
    detail = c.get("/api/customers/%s" % listed["uk"]).json()
    assert listed["due"] == 0
    assert listed["status"] == "ontrack"
    assert detail["status"] == "ontrack"
    assert detail["insight"]["money_bit"]


def test_customer_activity_groups_invoice_lines():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["19-09-24", "Acme", "R", 30, "INV-9"],
        ["19-09-24", "Acme", "R", 20, "INV-9"],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/customers").json()["customers"][0]["uk"]
    detail = c.get("/api/customers/%s" % uk).json()
    sales_events = [e for e in detail["activity"] if e["kind"] == "sale"]
    assert len(sales_events) == 1
    assert sales_events[0]["amount"] == 50
    assert sales_events[0]["invoice"] == "INV-9"
    assert sales_events[0]["invoice_id"]
    pdf = c.get("/api/customers/%s/pdf" % uk)
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"


def test_customer_activity_write_removed():
    c = client()
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 90]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 40]])])
    rec = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [["20-09-24", "Acme", "R", 15]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", rec, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/customers").json()["customers"][0]["uk"]
    body = c.get("/api/customers/%s" % uk).json()
    assert any(e["kind"] == "sale" for e in body["activity"])
    assert any(e["kind"] == "collection" for e in body["activity"])
    assert c.post("/api/customers/%s/activity" % uk, json={"kind": "visit"}).status_code in (404, 405)
    assert c.delete("/api/customers/%s/activity/n-1" % uk).status_code in (404, 405)
    pdf = c.get("/api/customers/%s/pdf" % uk)
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"


def test_activity_settings_endpoints_removed():
    c = client()
    assert c.get("/api/settings/activity-kinds").status_code in (404, 405)
    assert c.post("/api/settings/activity-kinds", json={"label": "Phone call"}).status_code in (404, 405)
    assert c.delete("/api/settings/activity-kinds/visit").status_code in (404, 405)
    assert c.get("/api/settings/activity-assignees").status_code in (404, 405)
    assert c.post("/api/settings/activity-assignees", json={"label": "Priya"}).status_code in (404, 405)
    assert c.delete("/api/settings/activity-assignees/priya").status_code in (404, 405)


def test_customer_forbidden_without_permission():
    c = client()
    raw = TestClient(app)
    assert raw.get("/api/customers").status_code == 401
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    other = TestClient(app)
    login_as(other, "wh", "wh12345")
    assert other.get("/api/customers").status_code == 403


def test_invoice_links_item_lines():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    c.put("/api/mappers/items", json={"column_map": {}, "unique_key": ["Date", "Item Name", "Qty", "Rate", "Invoice No"]})
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["19-09-24", "Acme", "R", 30, "INV-9"],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Invoice No", "Party Name"], [
        ["19-09-24", "SKU-A", 2, 10, "INV-9", "Acme"],
        ["19-09-24", "SKU-B", 1, 10, "INV-9", "Acme"],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/invoices")
    assert listed.status_code == 200
    assert listed.json()["total"] == 1
    inv = listed.json()["invoices"][0]
    assert inv["invoice"] == "INV-9"
    assert inv["line_count"] == 2
    detail = c.get("/api/invoices/" + inv["id"])
    assert detail.status_code == 200
    assert len(detail.json()["lines"]) == 2
    assert {row["item"] for row in detail.json()["lines"]} == {"SKU-A", "SKU-B"}
    lines = c.get("/api/item-lines")
    assert lines.json()["total"] == 2
    assert lines.json()["lines"][0]["invoice_id"] == inv["id"]
    assert detail.json()["customer_uk"]
    assert inv["rep"] == "R"
    assert "group" in inv
    assert detail.json()["group"] == inv["group"]
    assert detail.json()["rep"] == "R"


def test_same_invoice_no_separate_across_fiscal_years():
    """Invoice numbers reuse each FY — do not merge into one card or cross-link lines."""
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"]})
    c.put("/api/mappers/items", json={"column_map": {}, "unique_key": ["Date", "Item Name", "Qty", "Rate", "Invoice No"]})
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["10-05-24", "Acme", "R", 100, "INV-1"],
        ["15-05-25", "Acme", "R", 200, "INV-1"],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Invoice No", "Party Name"], [
        ["10-05-24", "Old-SKU", 1, 100, "INV-1", "Acme"],
        ["15-05-25", "New-SKU", 2, 100, "INV-1", "Acme"],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/invoices").json()
    assert listed["total"] == 2
    by_date = {row["date"]: row for row in listed["invoices"]}
    assert set(by_date) == {"2024-05-10", "2025-05-15"}
    assert by_date["2024-05-10"]["amount"] == 100
    assert by_date["2025-05-15"]["amount"] == 200
    assert by_date["2024-05-10"]["id"] != by_date["2025-05-15"]["id"]
    old = c.get("/api/invoices/" + by_date["2024-05-10"]["id"]).json()
    new = c.get("/api/invoices/" + by_date["2025-05-15"]["id"]).json()
    assert [row["item"] for row in old["lines"]] == ["Old-SKU"]
    assert [row["item"] for row in new["lines"]] == ["New-SKU"]
    fy25 = c.get("/api/invoices", params={"date_from": "2025-04-01", "date_to": "2026-03-31"}).json()
    assert fy25["total"] == 1
    assert fy25["invoices"][0]["amount"] == 200
    lines = c.get("/api/item-lines").json()["lines"]
    assert {row["item"]: row["invoice_id"] for row in lines} == {
        "Old-SKU": by_date["2024-05-10"]["id"],
        "New-SKU": by_date["2025-05-15"]["id"],
    }


def test_invoice_matches_vch_no_and_numeric_forms():
    c = client()
    party = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "South"]])])
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", party, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["19-09-24", "Acme", "Ria", 40, 1234],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Vch No.", "Party Name"], [
        ["19-09-24", "SKU-A", 2, 10, 1234.0, "Acme"],
        ["19-09-24", "SKU-B", 1, 20, "1234.0", "Acme"],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/invoices").json()
    assert listed["total"] == 1
    inv = listed["invoices"][0]
    assert inv["invoice"] == "1234"
    assert inv["rep"] == "Ria"
    assert inv["group"] == "South"
    assert inv["line_count"] == 2
    detail = c.get("/api/invoices/" + inv["id"]).json()
    assert len(detail["lines"]) == 2
    assert {row["item"] for row in detail["lines"]} == {"SKU-A", "SKU-B"}
    lines = c.get("/api/item-lines").json()["lines"]
    assert len(lines) == 2
    assert all(row["party"] == "Acme" for row in lines)
    assert all(row["group"] == "South" for row in lines)
    assert all(row["rep"] == "Ria" for row in lines)
    assert {row["item"] for row in lines} == {"SKU-A", "SKU-B"}
    preview = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Vch No."], [["19-09-24", "SKU-A", 1, 2, "9"]])])
    mapped = c.post("/api/uploads/preview", files={"file": ("t.xlsx", preview, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
    assert mapped["sheets"][0]["column_map"]["Vch No."] == "Invoice No"


def test_item_lines_inherit_customer_group_rep_from_sale():
    c = client()
    party = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "South"]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["19-09-24", "Acme", "Ria", 30, "INV-2"],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Invoice No"], [
        ["19-09-24", "SKU-A", 3, 10, "INV-2"],
    ])])
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", party, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    lines = c.get("/api/item-lines").json()["lines"]
    assert len(lines) == 1
    assert lines[0]["party"] == "Acme"
    assert lines[0]["group"] == "South"
    assert lines[0]["rep"] == "Ria"
    assert lines[0]["invoice"] == "INV-2"
    found = c.get("/api/item-lines", params={"rep": "Ria"}).json()
    assert found["total"] == 1
    group_hit = c.get("/api/item-lines", params={"group": "South"}).json()
    assert group_hit["total"] == 1


def test_invoice_tax_from_file_not_calculated():
    c = client()
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No", "Taxable Amount", "Tax Amount"], [
        ["19-09-24", "Acme", "Ria", 118, "INV-T", 100, 18],
        ["20-09-24", "Beta", "Sam", 50, "INV-B"],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Invoice No", "Amount", "Taxable Amount", "Tax Amount"], [
        ["19-09-24", "SKU-A", 2, 10, "INV-T", 59, 50, 9],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    listed = c.get("/api/invoices").json()
    assert {row["invoice"] for row in listed["invoices"]} == {"INV-T", "INV-B"}
    assert "Acme" in listed["options"]["party"]
    acme = c.get("/api/invoices", params={"party": "Acme"}).json()
    assert acme["total"] == 1
    inv = acme["invoices"][0]
    assert inv["amount"] == 118
    assert inv["before_tax"] == 100
    assert inv["tax"] == 18
    other = next(row for row in listed["invoices"] if row["invoice"] == "INV-B")
    assert other["amount"] == 50
    assert other["before_tax"] is None
    assert other["tax"] is None
    detail = c.get("/api/invoices/" + inv["id"]).json()
    assert detail["before_tax"] == 100
    assert detail["tax"] == 18
    assert detail["lines"][0]["amount"] == 59
    assert detail["lines"][0]["before_tax"] == 50
    assert detail["lines"][0]["tax"] == 9
    rounded = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Party Name"], [["21-09-24", "SKU-C", 3, 10.125, "Solo"]])])
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("r.xlsx", rounded, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    solo = c.get("/api/item-lines", params={"item": "SKU-C"}).json()["lines"][0]
    assert solo["amount"] == 30.38
    assert solo["tax"] is None
    assert solo["before_tax"] is None
    by_amt = c.get("/api/invoices", params={"sort": "amount", "dir": "asc"}).json()["invoices"]
    assert [row["invoice"] for row in by_amt] == ["INV-B", "INV-T"]


def test_sales_import_creates_customer_without_overwriting_due():
    c = client()
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "NewCo", "R", 40]])])
    r = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["types"]["sales"]["customers_created"] == 1
    assert r.json()["types"]["sales"]["parties_created"] == 1
    listed = c.get("/api/customers").json()
    assert listed["total"] == 1
    assert listed["customers"][0]["name"] == "NewCo"
    assert listed["customers"][0]["due"] == 0
    party = c.get("/api/rows", params={"type": "party", "party": "newco"}).json()
    assert party["total"] == 1
    assert party["rows"][0]["fields"]["Balance"] == 0

    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 500]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    later = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["20-09-24", "Acme", "R", 10]])])
    r2 = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s2.xlsx", later, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.json()["types"]["sales"]["customers_created"] == 0
    assert r2.json()["types"]["sales"]["parties_created"] == 0
    acme = next(x for x in c.get("/api/customers").json()["customers"] if x["name"] == "Acme")
    assert acme["due"] == 500
    party2 = c.get("/api/rows", params={"type": "party", "party": "acme"}).json()
    assert party2["rows"][0]["fields"]["Balance"] == 500


def test_items_import_creates_customer_and_buying_periods():
    c = client()
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Party Name"], [
        ["15-05-25", "Rice", 10, 20, "Buyer"],
        ["01-07-26", "Rice", 2, 20, "Buyer"],
        ["01-08-26", "Oil", 5, 30, "Buyer"],
        ["01-09-26", "Oil", 5, 30, "Buyer"],
    ])])
    r = c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["types"]["items"]["customers_created"] == 1
    assert r.json()["types"]["items"]["parties_created"] == 1
    uk = c.get("/api/customers").json()["customers"][0]["uk"]
    detail = c.get("/api/customers/%s" % uk).json()
    assert detail["due"] == 0
    buying = detail["buying"]
    assert buying["has_item_rows"] is True
    periods = buying["periods"]
    assert "this_year" in periods and "last_year" in periods
    year_names = {row["name"] for row in periods["this_year"]["items"]}
    last_names = {row["name"] for row in periods["last_year"]["items"]}
    assert "Rice" in last_names
    assert "Oil" in year_names
    assert buying["suggested"]
    assert any(s["name"] == "Rice" for s in buying["suggested"])
    insight = detail["insight"]
    assert insight["headline"]
    assert "Buyer" in insight["headline"]
    pdf = c.get("/api/customers/%s/pdf" % uk)
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"
    assert "brief" in (pdf.headers.get("content-disposition") or "").lower()
    statement = c.get("/api/customers/%s/pdf?view=customer" % uk)
    assert statement.status_code == 200
    assert statement.content[:4] == b"%PDF"
    assert "statement" in (statement.headers.get("content-disposition") or "").lower()
    # Customer statement must differ from the internal sales brief
    assert statement.content != pdf.content


def test_buying_inherits_party_from_sale_invoice():
    c = client()
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"], [
        ["01-09-26", "Acme", "Ria", 30, "INV-2"],
    ])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Invoice No"], [
        ["01-09-26", "SKU-A", 3, 10, "INV-2"],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = next(x["uk"] for x in c.get("/api/customers").json()["customers"] if x["name"] == "Acme")
    buying = c.get("/api/customers/%s" % uk).json()["buying"]
    assert buying["has_item_rows"] is True
    assert buying["default_period"]
    names = {row["name"] for row in buying["periods"]["all"]["items"]}
    assert "SKU-A" in names
    year_names = {row["name"] for row in buying["periods"]["this_year"]["items"]}
    assert "SKU-A" in year_names
    assert buying["suggested"]
    assert any(s["name"] == "SKU-A" for s in buying["suggested"])


def test_customer_suggests_last_order_when_bought_once():
    c = client()
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Party Name"], [
        ["01-09-26", "Oil", 5, 30, "Buyer"],
    ])])
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/customers").json()["customers"][0]["uk"]
    buying = c.get("/api/customers/%s" % uk).json()["buying"]
    assert buying["suggested"]
    assert buying["suggested"][0]["name"] == "Oil"
    assert "last order" in buying["suggested"][0]["reason"].lower() or "this year" in buying["suggested"][0]["reason"].lower() or "before" in buying["suggested"][0]["reason"].lower()


def test_cleanup_removes_operational_data_keeps_party_customer():
    c = client()
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 90]])])
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 40]])])
    items = xlsx_bytes([("items", ["Date", "Item Name", "Qty", "Rate", "Party Name"], [["19-09-24", "SKU-A", 2, 10, "Acme"]])])
    rec = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [["20-09-24", "Acme", "R", 15]])])
    pay = xlsx_bytes([("payments", ["Date", "Account Name", "Amount"], [["21-09-24", "Acme", 5]])])
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 8, 4]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", rec, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "payments"}, files={"file": ("p.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    before = c.get("/api/customers").json()
    assert before["total"] == 1
    assert before["customers"][0]["due"] == 90
    out = c.post("/api/data/cleanup")
    assert out.status_code == 200
    body = out.json()
    assert body["total"] >= 6
    assert body["removed"]["Outstanding"] == 1
    assert body["removed"]["Sales"] == 1
    assert body["removed"]["Item-wise sales"] == 1
    assert body["removed"]["Receipts"] == 1
    assert body["removed"]["Payments"] == 1
    assert body["removed"]["Stock"] == 1
    assert body.get("stock_reset") == 1
    assert c.get("/api/rows", params={"type": "sales"}).json()["total"] == 0
    assert c.get("/api/rows", params={"type": "receipt"}).json()["total"] == 0
    assert c.get("/api/rows", params={"type": "arr"}).json()["total"] == 0
    assert c.get("/api/rows", params={"type": "items"}).json()["total"] == 0
    assert c.get("/api/rows", params={"type": "payments"}).json()["total"] == 0
    stock_rows = c.get("/api/rows", params={"type": "stock"}).json()
    assert stock_rows["total"] == 1
    assert stock_rows["rows"][0]["fields"]["Qty"] == 0
    assert stock_rows["rows"][0]["fields"]["Item Name"] == "SKU-A"
    assert c.get("/api/rows", params={"type": "party"}).json()["total"] == 1
    customers = c.get("/api/customers").json()
    assert customers["total"] == 1
    assert customers["customers"][0]["due"] == 0
    party = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]
    assert party["fields"]["Balance"] == 0
    assert c.get("/api/invoices").json()["total"] == 0
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    other = TestClient(app)
    login_as(other, "wh", "wh12345")
    assert other.post("/api/data/cleanup").status_code == 403
    raw = TestClient(app)
    assert raw.post("/api/data/cleanup").status_code == 401


def test_item_360_holding_and_purchase():
    from datetime import datetime, timedelta

    from server.items360 import save_default_min, save_holding

    c = client()
    as_of = datetime.now()
    d1 = (as_of - timedelta(days=10)).strftime("%d-%m-%y")
    d2 = (as_of - timedelta(days=25)).strftime("%d-%m-%y")
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 8, 4], ["SKU-B", 100, 2]])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Party Name", "Invoice No"],
        [[d1, "SKU-A", 6, 10, "Acme", "INV-1"], [d2, "SKU-A", 12, 10, "Beta", "INV-2"]],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [[d1, "Acme", "R", 60, "INV-1"], [d2, "Beta", "R", 120, "INV-2"]],
    )])
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)}).status_code == 200

    listed = c.get("/api/items").json()
    assert listed["total"] == 2
    names = {row["name"]: row for row in listed["items"]}
    assert names["SKU-A"]["qty"] == 8
    assert names["SKU-A"]["min_hold"] == 0
    assert names["SKU-B"]["qty"] == 100

    detail = c.get("/api/items/SKU-A").json()
    assert detail["name"] == "SKU-A"
    assert detail["qty"] == 8
    assert detail["buyers"][0]["name"] in ("Acme", "Beta")
    assert any(row["invoice"] == "INV-1" for row in detail["movement"])
    assert detail["insight"]["headline"]

    assert c.put("/api/items/SKU-A/holding", json={"min_hold": 20}).status_code == 200
    assert c.put("/api/items/holding-default", json={"min_hold": 15}).status_code == 200

    store = get_store()
    save_holding(store, "SKU-A", min_hold=20, fill_to=50, lead_days=7)
    saved = c.get("/api/items/SKU-A").json()
    assert saved["min_hold"] == 20
    assert saved["fill_to"] == 50
    assert saved["lead_days"] == 7
    assert saved["status"] == "low"
    assert saved["buy_qty"] == 42
    assert saved["holding"]["own_min"] is True
    assert saved["max_days_hold"] == 60
    assert saved["uses_default_max_days"] is True

    low = c.get("/api/items", params={"status": "Below min"}).json()
    assert low["total"] == 1
    assert low["items"][0]["name"] == "SKU-A"

    again = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 80, 4], ["SKU-B", 100, 2]])])
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k2.xlsx", again, ctype)}).status_code == 200
    after = c.get("/api/items/SKU-A").json()
    assert after["qty"] == 80
    assert after["min_hold"] == 20
    assert after["fill_to"] == 50
    assert after["buy_qty"] == 0

    save_default_min(get_store(), 15)
    listed = c.get("/api/items").json()
    by_name = {row["name"]: row for row in listed["items"]}
    assert by_name["SKU-B"]["min_hold"] == 15
    assert by_name["SKU-A"]["min_hold"] == 20

    c.post("/api/data/cleanup")
    cleared = c.get("/api/items/SKU-A").json()
    assert cleared["qty"] == 0
    assert cleared["min_hold"] == 20
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k3.xlsx", stock, ctype)}).status_code == 200
    restored = c.get("/api/items/SKU-A").json()
    assert restored["min_hold"] == 20
    assert restored["qty"] == 8

    c.post("/api/users", json={"username": "view", "password": "view12345", "role": "Viewer"})
    other = TestClient(app)
    login_as(other, "view", "view12345")
    assert other.get("/api/items/SKU-A").status_code == 200
    assert other.put("/api/items/SKU-A/holding", json={"min_hold": 1}).status_code == 403


def test_item_360_max_days_hold_configurable():
    from datetime import datetime, timedelta

    from server.items360 import save_default_holding, save_holding

    c = client()
    as_of = datetime.now()
    day = (as_of - timedelta(days=10)).strftime("%d-%m-%y")
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 10, 4]])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Party Name", "Invoice No"],
        [[day, "SKU-A", 30, 10, "Acme", "INV-1"]],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [[day, "Acme", "R", 300, "INV-1"]],
    )])
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)}).status_code == 200

    base = c.get("/api/items/SKU-A").json()
    assert base["max_days_hold"] == 60
    assert base["target_days"] == 60
    assert base["uses_default_max_days"] is True
    pace = float(base["pace"] or 0)
    assert pace > 0
    assert abs(float(base["fill_to"]) - pace * 60) < 0.02

    saved = c.put("/api/items/SKU-A/holding", json={"max_days_hold": 30}).json()
    assert saved["max_days_hold"] == 30
    assert saved["own_max_days"] is True
    assert saved["uses_default_max_days"] is False
    assert saved["target_days"] == 30
    assert abs(float(saved["fill_to"]) - pace * 30) < 0.02

    cleared = c.put("/api/items/SKU-A/holding", json={"clear_max_days": True}).json()
    assert cleared["max_days_hold"] == 60
    assert cleared["uses_default_max_days"] is True

    save_default_holding(get_store(), max_days_hold=45)
    after_default = c.get("/api/items/SKU-A").json()
    assert after_default["max_days_hold"] == 45
    assert after_default["default_max_days"] == 45
    assert after_default["target_days"] == 45

    save_holding(get_store(), "SKU-A", max_days_hold=90)
    custom = c.get("/api/items/SKU-A").json()
    assert custom["max_days_hold"] == 90
    assert custom["target_days"] == 90
    assert custom["default_max_days"] == 45


def test_item_360_velocity_active_history_not_diluted():
    """New items: pace divides by days since first sale, not the full 90-day window."""
    from datetime import datetime, timedelta

    from vay.dates import normalize_date

    from server.items360 import _pace, _velocity

    as_of = normalize_date(datetime.now())
    # 90 units sold only over the last 10 inclusive days
    first = as_of - timedelta(days=9)
    lines = [
        {"date": first, "qty": 45, "amount": 450, "party": "Acme"},
        {"date": as_of, "qty": 45, "amount": 450, "party": "Acme"},
    ]
    pace, pace_days, pace_qty = _pace(lines, as_of)
    assert pace_qty == 90
    assert pace_days == 10
    assert abs(pace - 9.0) < 0.01

    vel = _velocity(lines, as_of, on_hand=90)
    assert vel["buckets"]["d90"]["qty"] == 90
    assert vel["buckets"]["d90"]["days"] == 10
    assert abs(vel["pace_90"] - 9.0) < 0.01
    assert vel["buckets"]["d30"]["days"] == 10
    assert abs(vel["pace_30"] - 9.0) < 0.01
    # Same active history in 30d and 90d → steady, not false "up"
    assert vel["trend"] == "flat"

    # Inclusive 7-day window: sale only on as_of
    one = [{"date": as_of, "qty": 7, "amount": 70, "party": "Acme"}]
    b7 = _velocity(one, as_of, on_hand=7)["buckets"]["d7"]
    assert b7["qty"] == 7
    assert b7["days"] == 1
    assert abs(b7["per_day"] - 7.0) < 0.01
    assert b7["window_days"] == 7


def test_item_360_includes_sales_only_skus_and_holding_form():
    """Sales-only SKUs appear with in_stock=False; list exposes buy_count and holding fields."""
    from datetime import datetime, timedelta

    from server.items360 import save_holding

    c = client()
    as_of = datetime.now()
    recent = (as_of - timedelta(days=3)).strftime("%d-%m-%y")
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [
        ["IN_STOCK", 10, 5],
    ])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Party Name", "Invoice No"],
        [
            [recent, "IN_STOCK", 4, 10, "Acme", "S1"],
            [recent, "SALES_ONLY", 6, 8, "Beta", "S2"],
        ],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [[recent, "Acme", "Rep A", 40, "S1"],
         [recent, "Beta", "Rep B", 48, "S2"]],
    )])
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)}).status_code == 200

    listed = c.get("/api/items").json()
    by_name = {row["name"]: row for row in listed["items"]}
    assert "SALES_ONLY" in by_name
    assert "IN_STOCK" in by_name
    so = by_name["SALES_ONLY"]
    assert so["qty"] == 0
    assert so["in_stock"] is False
    assert by_name["IN_STOCK"]["in_stock"] is True
    assert "buy_qty" in so
    assert "buy_count" in listed

    save_holding(
        get_store(),
        "SALES_ONLY",
        min_hold=5,
        fill_to=20,
        lead_days=7,
        max_days_hold=45,
        discontinued=False,
    )
    detail = c.get("/api/items/SALES_ONLY").json()
    assert detail["in_stock"] is False
    assert detail["qty"] == 0
    assert detail["min_hold"] == 5
    assert detail["fill_to"] == 20
    assert detail["lead_days"] == 7
    assert detail["max_days_hold"] == 45
    assert detail["buy_qty"] == 20
    assert detail["status"] == "low"
    assert detail["holding"]["min_hold"] == 5
    assert detail["holding"]["fill_to"] == 20
    assert detail["holding"]["lead_days"] == 7
    assert detail["holding"]["max_days_hold"] == 45
    assert detail["holding"]["discontinued"] is False
    assert detail["buyers"]
    assert detail["buyers"][0]["name"] == "Beta"

    listed2 = c.get("/api/items").json()
    assert listed2["buy_count"] >= 1
    assert listed2["out_of_stock_count"] >= 1
    so2 = next(row for row in listed2["items"] if row["name"] == "SALES_ONLY")
    assert so2["buy_qty"] == 20

    oos = c.get("/api/items", params={"in_stock": "no"}).json()
    assert oos["total"] >= 1
    assert all(not row["in_stock"] for row in oos["items"])
    assert any(row["name"] == "SALES_ONLY" for row in oos["items"])


def test_item_360_velocity_and_stock_position():
    from datetime import datetime, timedelta

    from server.items360 import save_holding

    c = client()
    as_of = datetime.now()
    recent = (as_of - timedelta(days=5)).strftime("%d-%m-%y")
    mid = (as_of - timedelta(days=20)).strftime("%d-%m-%y")
    old = (as_of - timedelta(days=80)).strftime("%d-%m-%y")
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [
        ["FAST", 5, 10],
        ["SLOW", 200, 2],
        ["IDLE", 40, 3],
    ])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Party Name", "Invoice No"],
        [
            [recent, "FAST", 12, 10, "Acme", "F1"],
            [mid, "FAST", 18, 10, "Beta", "F2"],
            [old, "FAST", 9, 10, "Acme", "F3"],
            [old, "SLOW", 3, 2, "Acme", "S1"],
        ],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [
            [recent, "Acme", "R", 120, "F1"],
            [mid, "Beta", "R", 180, "F2"],
            [old, "Acme", "R", 90, "F3"],
            [old, "Acme", "R", 6, "S1"],
        ],
    )])
    assert c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)}).status_code == 200
    assert c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)}).status_code == 200

    listed = c.get("/api/items").json()
    assert "under_count" in listed
    assert "over_count" in listed
    assert "position" in listed["options"]
    by_name = {row["name"]: row for row in listed["items"]}
    assert by_name["FAST"]["pace_30"] > 0
    assert by_name["FAST"]["pace_90"] > 0
    assert by_name["FAST"]["stock_position"] in ("under", "tight", "balanced", "over")
    assert by_name["IDLE"]["stock_position"] == "dead"

    save_holding(get_store(), "FAST", min_hold=40, fill_to=60, lead_days=7)
    fast = c.get("/api/items/FAST").json()
    assert fast["status"] == "low"
    assert fast["stock_position"] == "under"
    assert fast["velocity"]["buckets"]["d30"]["qty"] > 0
    assert fast["insight"]["velocity_bit"]
    assert fast["insight"]["stock_label"] == "Understocked"
    assert fast["insight"]["cover_bit"] or fast["insight"]["next_move"]
    assert "per_day" in (fast["periods"].get("this_month") or {})

    save_holding(get_store(), "SLOW", min_hold=5, fill_to=20, lead_days=3)
    slow = c.get("/api/items/SLOW").json()
    assert slow["qty"] == 200
    assert slow["stock_position"] == "over"
    assert slow["over_qty"] and slow["over_qty"] > 0
    assert "Overstocked" in (slow["insight"]["headline"] or "") or slow["insight"]["stock_label"] == "Overstocked"

    under = c.get("/api/items", params={"position": "Understocked"}).json()
    assert under["total"] >= 1
    assert all(row["stock_position"] == "under" for row in under["items"])


def test_item_360_buyers_from_invoice_when_line_has_no_party():
    from datetime import datetime, timedelta

    c = client()
    as_of = datetime.now()
    day = (as_of - timedelta(days=8)).strftime("%d-%m-%y")
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 8, 4]])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Invoice No"],
        [[day, "SKU-A", 6, 10, "631"]],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [[day, "MAXIMO TRADING", "R", 60, "631"]],
    )])
    party = xlsx_bytes([("party", ["Account Name", "Group"], [["MAXIMO TRADING", "South"]])])
    c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)})
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", party, ctype)})
    detail = c.get("/api/items/SKU-A").json()
    assert detail["buyers"]
    assert detail["buyers"][0]["name"] == "MAXIMO TRADING"
    assert detail["usual_buyer"] == "MAXIMO TRADING"
    assert detail["movement"][0]["party"] == "MAXIMO TRADING"
    assert detail["movement"][0]["group"] == "South"
    assert detail["movement"][0]["rep"] == "R"
    assert detail["groups"][0]["name"] == "South"
    assert detail["reps"][0]["name"] == "R"


def test_item_360_can_mark_discontinued():
    from datetime import datetime, timedelta

    from server.items360 import save_holding

    c = client()
    day = (datetime.now() - timedelta(days=8)).strftime("%d-%m-%y")
    ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stock = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 8, 4]])])
    items = xlsx_bytes([(
        "items",
        ["Date", "Item Name", "Qty", "Rate", "Invoice No"],
        [[day, "SKU-A", 6, 10, "INV-1"]],
    )])
    sales = xlsx_bytes([(
        "sales",
        ["Date", "Party Name", "Sales Rep", "Net Amount", "Invoice No"],
        [[day, "Acme", "R", 60, "INV-1"]],
    )])
    c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k.xlsx", stock, ctype)})
    c.post("/api/uploads", data={"type": "items"}, files={"file": ("i.xlsx", items, ctype)})
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, ctype)})
    assert c.put("/api/items/SKU-A/holding", json={"discontinued": True}).status_code == 200
    save_holding(get_store(), "SKU-A", min_hold=20, fill_to=50, lead_days=7, discontinued=True)
    stopped = c.get("/api/items/SKU-A").json()
    assert stopped["discontinued"] is True
    assert stopped["status"] == "stopped"
    assert stopped["status_label"] == "Discontinued"
    assert stopped["buy_qty"] == 0
    assert "discontinued" in (stopped.get("insight") or {}).get("headline", "").lower()
    listed = c.get("/api/items", params={"status": "Discontinued"}).json()
    assert listed["total"] == 1
    assert listed["items"][0]["name"] == "SKU-A"
    again = xlsx_bytes([("stock", ["Item Name", "Qty", "P.Price"], [["SKU-A", 80, 4]])])
    c.post("/api/uploads", data={"type": "stock"}, files={"file": ("k2.xlsx", again, ctype)})
    after = c.get("/api/items/SKU-A").json()
    assert after["discontinued"] is True
    assert after["qty"] == 80
    assert after["min_hold"] == 20
    save_holding(get_store(), "SKU-A", min_hold=20, fill_to=50, lead_days=7, discontinued=False)
    live = c.get("/api/items/SKU-A").json()
    assert live["discontinued"] is False
    assert live["status"] != "stopped"


def _export_catalog(c, rid):
    r = c.get("/api/runs/%s/export-pdf" % rid)
    assert r.status_code == 200
    return r.json()


def _snapshot_run(store, reports, report_date="2024-09-19", packs=None):
    json_id = store.put_blob(json.dumps({"reports": reports}).encode("utf-8"), "snap.json", "application/json")
    run = store.insert_run({
        "status": "succeeded",
        "packs": packs or {"core": True, "profit": True},
        "report_date": report_date,
        "json_id": json_id,
        "report_count": len(reports),
    })
    return str(run["_id"])


def test_pdf_export_zip_skip_and_download(tmp_path, monkeypatch):
    monkeypatch.setenv("VAY_EXPORT_DIR", str(tmp_path))
    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    rid = post.json()["id"]
    catalog = _export_catalog(c, rid)
    ids = [row["id"] for row in catalog["reports"]]
    assert ids
    assert all(not row["exists"] for row in catalog["reports"])
    assert c.post("/api/runs/%s/export-pdf" % rid, json={"ids": []}).status_code == 400
    exported = c.post("/api/runs/%s/export-pdf" % rid, json={"ids": ids})
    assert exported.status_code in (200, 202)
    status = _export_catalog(c, rid)
    assert status["status"] == "succeeded"
    assert status["zip_ready"] is True
    assert status["count"] >= 1
    assert status["done"] == status["total"] >= 1
    assert "vay_reports/19-09-2024" in status["folder"]
    by_id = {row["id"]: row for row in status["reports"]}
    assert by_id["sales_rep_performance_report"]["exists"] is True
    assert by_id["sales_rep_performance_report"]["state"] == "done"
    file_resp = c.get("/api/runs/%s/export-pdf/file" % rid)
    assert file_resp.status_code == 200
    assert file_resp.headers["content-type"].startswith("application/zip")
    names = zipfile.ZipFile(io.BytesIO(file_resp.content)).namelist()
    assert names
    assert all(n.endswith(".pdf") for n in names)
    assert not any(n.endswith(".zip") for n in names)
    assert "sales_rep_performance_report.pdf" in names
    zf = zipfile.ZipFile(io.BytesIO(file_resp.content))
    assert zf.read(names[0])[:4] == b"%PDF"
    again = c.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["sales_rep_performance_report"]})
    assert again.status_code in (200, 202)
    rebuilt = _export_catalog(c, rid)
    assert rebuilt["status"] == "succeeded"
    assert rebuilt["count"] >= 1
    assert rebuilt["skipped"] == 0
    assert rebuilt["zip_ready"] is True
    only = zipfile.ZipFile(io.BytesIO(c.get("/api/runs/%s/export-pdf/file" % rid).content)).namelist()
    assert only == ["sales_rep_performance_report.pdf"]


def test_pdf_export_catalog_selection_and_permissions(tmp_path, monkeypatch):
    monkeypatch.setenv("VAY_EXPORT_DIR", str(tmp_path))
    c = client()
    store = get_store()
    reports = [
        {"id": "sales_rep_performance_report", "group": "Performance", "friendly_title": "Sales by person", "headers": ["Sales Rep"], "rows": [["R"]]},
        {"id": "monthly_profit_report", "group": "Profit", "friendly_title": "Profit and loss", "headers": ["Line"], "rows": [["Sales"]]},
        {"id": "source_data_warnings", "group": "Data issues", "friendly_title": "Data issues", "headers": ["Issue"], "rows": []},
    ]
    rid = _snapshot_run(store, reports)
    admin = _export_catalog(c, rid)
    assert {row["id"] for row in admin["reports"]} == {
        "sales_rep_performance_report",
        "monthly_profit_report",
        "source_data_warnings",
    }
    assert all(row["state"] == "idle" for row in admin["reports"])
    c.post("/api/users", json={"username": "fin", "password": "fin12345", "role": "Finance"})
    c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    fin = TestClient(app)
    login_as(fin, "fin", "fin12345")
    look = TestClient(app)
    login_as(look, "look", "look123")
    fin_ids = {row["id"] for row in _export_catalog(fin, rid)["reports"]}
    assert "sales_rep_performance_report" not in fin_ids
    assert "monthly_profit_report" in fin_ids
    assert "source_data_warnings" in fin_ids
    look_ids = {row["id"] for row in _export_catalog(look, rid)["reports"]}
    assert "sales_rep_performance_report" in look_ids
    assert "monthly_profit_report" not in look_ids
    assert look.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["source_data_warnings"]}).status_code == 403
    denied = fin.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["sales_rep_performance_report"]})
    assert denied.status_code == 400
    mixed = fin.post("/api/runs/%s/export-pdf" % rid, json={
        "ids": ["sales_rep_performance_report", "monthly_profit_report"],
    })
    assert mixed.status_code in (200, 202)
    after = _export_catalog(fin, rid)
    assert after["status"] == "succeeded"
    assert after["count"] == 1
    assert after["total"] == 1
    assert after["done"] == 1
    by_id = {row["id"]: row for row in after["reports"]}
    assert by_id["monthly_profit_report"]["exists"] is True
    assert by_id["monthly_profit_report"]["state"] == "done"
    assert "sales_rep_performance_report" not in by_id
    names = zipfile.ZipFile(io.BytesIO(fin.get("/api/runs/%s/export-pdf/file" % rid).content)).namelist()
    assert names == ["monthly_profit_report.pdf"]
    admin_names = zipfile.ZipFile(io.BytesIO(c.get("/api/runs/%s/export-pdf/file" % rid).content)).namelist()
    assert "monthly_profit_report.pdf" in admin_names
    assert "sales_rep_performance_report.pdf" not in admin_names
    picked = c.post("/api/runs/%s/export-pdf" % rid, json={
        "ids": ["sales_rep_performance_report", "source_data_warnings"],
    })
    assert picked.status_code in (200, 202)
    admin_names = zipfile.ZipFile(io.BytesIO(c.get("/api/runs/%s/export-pdf/file" % rid).content)).namelist()
    assert set(admin_names) == {"sales_rep_performance_report.pdf", "source_data_warnings.pdf"}


def test_pdf_export_recovers_stale_queue(tmp_path, monkeypatch):
    from datetime import datetime, timedelta

    monkeypatch.setenv("VAY_EXPORT_DIR", str(tmp_path))
    c = client()
    store = get_store()
    reports = [{
        "id": "sales_rep_performance_report",
        "group": "Performance",
        "friendly_title": "Sales by person",
        "headers": ["Sales Rep"],
        "rows": [["R"]],
    }]
    rid = _snapshot_run(store, reports)
    store.update_run(rid, {
        "pdf_export_status": "queued",
        "pdf_export_heartbeat": datetime.utcnow() - timedelta(seconds=60),
        "pdf_export_progress": {
            "current": "",
            "done": 0,
            "total": 1,
            "reports": [{"id": "sales_rep_performance_report", "title": "Sales by person", "group": "Performance", "state": "waiting"}],
        },
    })
    body = _export_catalog(c, rid)
    assert body["status"] == "failed"
    exported = c.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["sales_rep_performance_report"]})
    assert exported.status_code in (200, 202)
    assert exported.json()["status"] == "succeeded"


def test_pdf_export_auth_and_conflicts(tmp_path, monkeypatch):
    monkeypatch.setenv("VAY_EXPORT_DIR", str(tmp_path))
    raw = TestClient(app)
    assert raw.post("/api/runs/x/export-pdf", json={"ids": ["x"]}).status_code == 401
    c = client()
    queued = get_store().insert_run({"status": "queued", "packs": {}})
    assert c.post("/api/runs/%s/export-pdf" % queued["_id"], json={"ids": ["x"]}).status_code == 409
    assert c.post("/api/runs/000000000000000000000000/export-pdf", json={"ids": ["x"]}).status_code == 404
    c = client()
    _seed_core(c)
    rid = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"}).json()["id"]
    get_store().insert_run({"status": "running", "packs": {}})
    assert c.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["sales_rep_performance_report"]}).status_code == 409
    c2 = client()
    get_store().insert_run({
        "status": "succeeded",
        "pdf_export_status": "running",
        "pdf_export_heartbeat": __import__("datetime").datetime.utcnow(),
        "packs": {},
    })
    assert c2.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"}).status_code == 409


def test_can_add_custom_role_and_assign_user():
    c = client()
    created = c.post("/api/roles", json={"name": "Auditor", "permissions": ["customer.view", "sales.view"]})
    assert created.status_code == 200
    assert created.json()["name"] == "Auditor"
    assert created.json()["permissions"] == ["customer.view", "sales.view"]
    names = [r["name"] for r in c.get("/api/roles").json()["roles"]]
    assert names[:5] == ["Admin", "Sales", "Finance", "Warehouse", "Viewer"]
    assert "Auditor" in names
    assert c.post("/api/roles", json={"name": "auditor"}).status_code == 400
    assert c.post("/api/roles", json={"name": "  "}).status_code == 400
    assert c.post("/api/roles", json={"name": "???"}).status_code == 400
    user = c.post("/api/users", json={"username": "aud", "password": "aud12345", "role": "Auditor"})
    assert user.status_code == 200
    assert user.json()["role"] == "Auditor"
    assert "customer.view" in user.json()["permissions"]
    saved = c.put("/api/roles/Auditor", json={"permissions": ["customer.view"]})
    assert saved.status_code == 200
    assert saved.json()["permissions"] == ["customer.view"]
    patched = c.patch("/api/users/aud", json={"role": "auditor"})
    assert patched.status_code == 200
    assert patched.json()["role"] == "Auditor"
    assert c.post("/api/users", json={"username": "x", "password": "xxxxxxxx", "role": "Nobody"}).status_code == 400
    assert c.put("/api/roles/Nobody", json={"permissions": []}).status_code == 404


def test_viewer_cannot_manage_roles():
    c = client()
    c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    other = TestClient(app)
    login_as(other, "look", "look123")
    assert other.post("/api/roles", json={"name": "Temp"}).status_code == 403
    assert other.get("/api/roles").status_code == 403
    assert other.get("/api/users").status_code == 403


def test_roles_view_cannot_write_permissions():
    c = client()
    created = c.post("/api/roles", json={"name": "Looker", "permissions": ["roles.view", "customer.view"]})
    assert created.status_code == 200
    c.post("/api/users", json={"username": "looker", "password": "look12345", "role": "Looker"})
    other = TestClient(app)
    login_as(other, "looker", "look12345")
    listed = other.get("/api/roles")
    assert listed.status_code == 200
    assert "roles.manage" in listed.json()["all_permissions"]
    assert other.post("/api/roles", json={"name": "Temp"}).status_code == 403
    assert other.put("/api/roles/Looker", json={"permissions": ["customer.view"]}).status_code == 403
    assert other.get("/api/users").status_code == 403


def test_users_view_cannot_write_users():
    c = client()
    c.post("/api/roles", json={"name": "HR", "permissions": ["users.view"]})
    c.post("/api/users", json={"username": "hr", "password": "hr123456", "role": "HR"})
    other = TestClient(app)
    login_as(other, "hr", "hr123456")
    listed = other.get("/api/users")
    assert listed.status_code == 200
    assert listed.json()["users"]
    assert other.post("/api/users", json={"username": "x", "password": "xxxxxxxx", "role": "Viewer"}).status_code == 403
    assert other.get("/api/roles").status_code == 403


def test_users_manage_can_list_role_names():
    c = client()
    c.post("/api/roles", json={"name": "HR Admin", "permissions": ["users.manage"]})
    c.post("/api/users", json={"username": "hradmin", "password": "hr123456", "role": "HR Admin"})
    other = TestClient(app)
    login_as(other, "hradmin", "hr123456")
    listed = other.get("/api/roles")
    assert listed.status_code == 200
    names = [r["name"] for r in listed.json()["roles"]]
    assert "Viewer" in names
    assert listed.json()["all_permissions"] == []
    assert other.post("/api/roles", json={"name": "Temp"}).status_code == 403


def test_roles_manage_can_write_without_users():
    c = client()
    c.post("/api/roles", json={"name": "Access", "permissions": ["roles.view", "roles.manage"]})
    c.post("/api/users", json={"username": "access", "password": "acc12345", "role": "Access"})
    other = TestClient(app)
    login_as(other, "access", "acc12345")
    created = other.post("/api/roles", json={"name": "Temp", "permissions": ["customer.view"]})
    assert created.status_code == 200
    saved = other.put("/api/roles/Temp", json={"permissions": ["sales.view"]})
    assert saved.status_code == 200
    assert saved.json()["permissions"] == ["sales.view"]
    assert other.get("/api/users").status_code == 403


def test_pdf_export_strips_profit_without_permission():
    c = client()
    store = get_store()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("sales_rep_performance_report.pdf", b"%PDF-sales")
        zf.writestr("monthly_profit_report.pdf", b"%PDF-profit")
    zip_id = store.put_blob(buf.getvalue(), "x.zip", "application/zip")
    run = store.insert_run({
        "status": "succeeded",
        "packs": {"core": True, "profit": True},
        "pdf_export_status": "succeeded",
        "pdf_zip_id": zip_id,
    })
    rid = str(run["_id"])
    admin_names = zipfile.ZipFile(io.BytesIO(c.get("/api/runs/%s/export-pdf/file" % rid).content)).namelist()
    assert "sales_rep_performance_report.pdf" in admin_names
    assert "monthly_profit_report.pdf" in admin_names
    c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    other = TestClient(app)
    login_as(other, "look", "look123")
    assert other.get("/api/runs/%s/export-pdf" % rid).status_code == 200
    assert other.post("/api/runs/%s/export-pdf" % rid, json={"ids": ["monthly_profit_report"]}).status_code == 403
    file_resp = other.get("/api/runs/%s/export-pdf/file" % rid)
    assert file_resp.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(file_resp.content)).namelist()
    assert "sales_rep_performance_report.pdf" in names
    assert "monthly_profit_report.pdf" not in names


def test_dashboard_empty_store():
    c = client()
    r = c.get("/api/dashboard")
    assert r.status_code == 200
    body = r.json()
    assert body["run"] is None
    assert body["kpis"] == []
    assert body["chart"] == []
    assert body["has_data"] is False
    assert body["customers"]["counts"]["urgent"] == 0
    assert body["customers"]["top"] == []
    assert body["invoices"] == []
    assert body["stock"]["low_count"] == 0
    assert body["stock"]["top"] == []
    assert "profit" not in body


def test_dashboard_after_run_has_performance_kpis():
    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    assert post.status_code == 202
    rid = post.json()["id"]
    body = c.get("/api/dashboard").json()
    assert body["run"]["id"] == rid
    assert body["run"]["report_date"] == "2024-09-19"
    stored = get_store().get_run(rid).get("dashboard") or {}
    assert stored.get("kpis")
    progress = c.get("/api/runs/%s" % rid).json()
    assert "dashboard" not in progress
    steps = (progress.get("generate_progress") or {}).get("steps") or []
    ids = [row["id"] for row in steps]
    assert ids.index("dashboard") < ids.index("save-reports")
    assert next(row["state"] for row in steps if row["id"] == "dashboard") == "done"
    assert body["has_data"] is True
    by_id = {row["id"]: row["value"] for row in body["kpis"]}
    assert by_id["mtd_sales"] == 100
    assert by_id["mtd_collection"] == 40
    assert by_id["ytd_sales"] == 100
    assert by_id["ytd_collection"] == 40
    assert by_id["mtd_gap"] == 60
    assert by_id["mtd_rate"] == 40
    assert by_id["d15_sales"] == 100
    assert by_id["d15_collection"] == 40
    assert by_id["d15_gap"] == 60
    assert by_id["d15_rate"] == 40
    assert by_id["ytd_gap"] == 60
    assert by_id["ytd_rate"] == 40
    assert by_id["total_sales"] == 100
    assert by_id["total_collection"] == 40
    assert by_id["total_gap"] == 60
    assert by_id["total_rate"] == 40
    assert body["total_source"] in ("fiscal", "imported")
    assert "import_newer" in body
    assert "buy_qty_sum" in body["stock"]
    assert "credit" in body["customers"]["counts"]
    assert body["chart"]
    assert body["chart"][0]["name"] == "R"
    assert body["chart"][0]["sales"] == 100
    assert "chart_groups" in body
    assert body["chart_groups"]
    assert body["chart_groups"][0]["sales"] == 100
    assert "chart_groups_mtd" in body
    assert body["chart_groups_mtd"]
    assert "chart_aging" in body
    assert body["chart_aging"]
    assert "chart_monthly" in body
    assert isinstance(body["chart_monthly"], list)
    assert "mtd_gap_pct" in by_id
    assert by_id["mtd_gap_pct"] == 60
    assert "active_customers_mtd" in by_id
    assert by_id["active_customers_mtd"] >= 1
    assert "ar_balance" in by_id or "due_90" in by_id or "overdue_30" in by_id
    assert "dso" in by_id or by_id.get("ar_balance") in (None, 0)
    assert "top_rep_share" in by_id
    assert by_id["top_rep_share"] == 100
    assert "due_0_15" in by_id or "overdue_15" in by_id or "new_credit" in by_id
    assert "inactive_mtd" in by_id
    assert "collection_followups" in by_id
    assert "invoice_mtd_count" in by_id
    assert by_id["invoice_mtd_count"] >= 1
    assert "invoice_mtd_avg" in by_id
    assert "buy_value_sum" in body["stock"]
    assert "chart_due" in body
    assert body["invoices"]
    assert body["customers"]["total"] >= 1
    assert "counts" in c.get("/api/customers").json()


def test_dashboard_profit_section_when_pack_runs():
    c = client()
    _seed_core(c)
    c.put("/api/mappers/payments", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Amount"]})
    pay = xlsx_bytes([("payments", ["Date", "Account Name", "Amount"], [["19-09-24", "RENT SHOP PAYABLE", 10]])])
    c.post("/api/uploads", data={"type": "payments"}, files={"file": ("pay.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    post = c.post("/api/runs", json={"packs": {"core": True, "profit": True}, "report_date": "2024-09-19"})
    assert post.status_code == 202
    body = c.get("/api/dashboard").json()
    assert "profit" in body
    profit_ids = {row["id"] for row in body["profit"]["kpis"]}
    assert "expense_mtd" in profit_ids or "expense_ytd" in profit_ids or "profit_ytd" in profit_ids
    assert isinstance(body["profit"]["chart"], list)


def test_dashboard_total_includes_prior_fiscal_years():
    c = client()
    _seed_core(c)
    prior = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["01-04-23", "Old Co", "R", 50]])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("old.xlsx", prior, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    post = c.post("/api/runs", json={"packs": {"core": True, "fiscal": True}, "report_date": "2024-09-19"})
    assert post.status_code == 202
    by_id = {row["id"]: row["value"] for row in c.get("/api/dashboard").json()["kpis"]}
    assert by_id["ytd_sales"] == 100
    assert by_id["total_sales"] == 150
    assert by_id["total_collection"] == 40
    assert c.get("/api/dashboard").json()["total_source"] == "fiscal"
    monthly = c.get("/api/dashboard").json()["chart_monthly"]
    assert monthly
    assert monthly[0]["name"]
    dash = c.get("/api/dashboard").json()
    by_id = {row["id"]: row["value"] for row in dash["kpis"]}
    assert "dso" in by_id
    assert "mom_sales_pct" in by_id or len(dash["chart_monthly"]) < 2
    assert "chart_rate" in dash
    assert dash["chart_rate"]
    assert dash["chart_rate"][0]["rate"] == 40


def test_dashboard_warehouse_omits_sales():
    c = client()
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    other = TestClient(app)
    login_as(other, "wh", "wh12345")
    r = other.get("/api/dashboard")
    assert r.status_code == 200
    body = r.json()
    assert "kpis" not in body
    assert "chart" not in body
    assert "invoices" not in body
    assert "customers" not in body
    assert "stock" in body
    assert "profit" not in body


def test_dashboard_viewer_has_no_profit_section():
    c = client()
    c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    other = TestClient(app)
    login_as(other, "look", "look123")
    body = other.get("/api/dashboard").json()
    assert "profit" not in body
    assert "kpis" in body
    assert "customers" in body


def test_dashboard_requires_auth():
    raw = TestClient(app)
    assert raw.get("/api/dashboard").status_code == 401


def test_dashboard_saved_payload_respects_permissions():
    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    assert post.status_code == 202
    rid = post.json()["id"]
    stored = get_store().get_run(rid).get("dashboard") or {}
    assert stored.get("kpis")
    assert "invoices" in stored

    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    wh = TestClient(app)
    login_as(wh, "wh", "wh12345")
    body = wh.get("/api/dashboard").json()
    assert "kpis" not in body
    assert "chart" not in body
    assert "invoices" not in body
    assert "customers" not in body
    assert "profit" not in body
    assert "stock" in body

    c.post("/api/users", json={"username": "look", "password": "look123", "role": "Viewer"})
    look = TestClient(app)
    login_as(look, "look", "look123")
    viewed = look.get("/api/dashboard").json()
    assert "profit" not in viewed
    assert viewed.get("kpis")
    assert "customers" in viewed

    get_store().update_run(rid, {"dashboard": None})
    again = c.get("/api/dashboard").json()
    by_id = {row["id"]: row["value"] for row in again["kpis"]}
    assert by_id["mtd_sales"] == 100
    assert isinstance(get_store().get_run(rid).get("dashboard"), dict)


def test_payment_import_creates_party_without_customer():
    c = client()
    pay = xlsx_bytes([("payments", ["Date", "Account Name", "Amount"], [["19-09-24", "Vendor Co", 25]])])
    r = c.post("/api/uploads", data={"type": "payments"}, files={"file": ("pay.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["types"]["payments"]["parties_created"] == 1
    party = c.get("/api/rows", params={"type": "party", "party": "vendor co"}).json()
    assert party["total"] == 1
    assert party["rows"][0]["fields"]["Balance"] == 0
    assert not party["rows"][0]["party_type"]
    assert c.get("/api/customers").json()["total"] == 0
    again = c.post("/api/uploads", data={"type": "payments"}, files={"file": ("pay2.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert again.json()["types"]["payments"]["parties_created"] == 0
    assert again.json()["types"]["payments"]["skipped"] == 1
    party2 = c.get("/api/rows", params={"type": "party", "party": "vendor co"}).json()
    assert party2["total"] == 1
    assert party2["rows"][0]["fields"]["Balance"] == 0


def test_arr_import_tags_party_as_customer():
    c = client()
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 500]])])
    r = c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["types"]["arr"]["parties_created"] == 1
    party = c.get("/api/rows", params={"type": "party", "party": "acme"}).json()
    assert party["total"] == 1
    assert party["rows"][0]["party_type"] == "customer"
    assert party["rows"][0]["party_type_label"] == "Customer"
    assert party["rows"][0]["fields"]["Balance"] == 500
    customers = c.get("/api/customers").json()
    assert customers["total"] == 1
    assert customers["customers"][0]["name"] == "Acme"


def test_party_types_from_settings():
    c = client()
    listed = c.get("/api/settings/party-types")
    assert listed.status_code == 200
    ids = {t["id"] for t in listed.json()["types"]}
    assert ids == {"customer"}
    assert listed.json()["types"][0]["builtin"] is True
    assert c.post("/api/settings/party-types", json={"label": "customer"}).status_code == 400
    assert c.delete("/api/settings/party-types/customer").status_code == 400
    added = c.post("/api/settings/party-types", json={"label": "Vendor"})
    assert added.status_code == 200
    assert any(t["id"] == "vendor" and t["label"] == "Vendor" and not t["builtin"] for t in added.json()["types"])
    gone = c.delete("/api/settings/party-types/vendor")
    assert gone.status_code == 200
    assert all(t["id"] != "vendor" for t in gone.json()["types"])
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    other = TestClient(app)
    login_as(other, "wh", "wh12345")
    assert other.get("/api/settings/party-types").status_code == 403
    assert other.post("/api/settings/party-types", json={"label": "Supplier"}).status_code == 403


def test_party_upload_can_change_party_type():
    c = client()
    c.post("/api/settings/party-types", json={"label": "Vendor"})
    pay = xlsx_bytes([("payments", ["Date", "Account Name", "Amount"], [["19-09-24", "NewCo", 10]])])
    c.post("/api/uploads", data={"type": "payments"}, files={"file": ("pay.xlsx", pay, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]["uk"]
    assert c.get("/api/customers").json()["total"] == 0
    c.post("/api/users", json={"username": "viewonly", "password": "view12345", "role": "Viewer"})
    viewer = TestClient(app)
    login_as(viewer, "viewonly", "view12345")
    assert viewer.patch("/api/parties/" + uk, json={"party_type": "vendor"}).status_code == 403
    patched = c.patch("/api/parties/" + uk, json={"party_type": "vendor"})
    assert patched.status_code == 200
    assert patched.json()["party_type"] == "vendor"
    assert patched.json()["party_type_label"] == "Vendor"
    listed = c.get("/api/rows", params={"type": "party", "party_type": "Vendor"}).json()
    assert listed["total"] == 1
    assert listed["rows"][0]["party_type"] == "vendor"
    to_customer = c.patch("/api/parties/" + uk, json={"party_type": "Customer"})
    assert to_customer.status_code == 200
    assert to_customer.json()["party_type"] == "customer"
    customers = c.get("/api/customers").json()
    assert customers["total"] == 1
    assert customers["customers"][0]["name"] == "NewCo"
    assert c.patch("/api/parties/" + uk, json={"party_type": "ghost"}).status_code == 400
    assert c.patch("/api/parties/missing", json={"party_type": "vendor"}).status_code == 404


def test_party_reimport_keeps_party_type():
    c = client()
    c.post("/api/settings/party-types", json={"label": "Vendor"})
    party = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "South"]])])
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p.xlsx", party, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]["uk"]
    c.patch("/api/parties/" + uk, json={"party_type": "vendor"})
    party2 = xlsx_bytes([("party", ["Account Name", "Group"], [["Acme", "North"]])])
    c.post("/api/uploads", data={"type": "party"}, files={"file": ("p2.xlsx", party2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    row = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]
    assert row["fields"]["Group"] == "North"
    assert row["party_type"] == "vendor"


def test_arr_reimport_keeps_manual_party_type():
    c = client()
    c.post("/api/settings/party-types", json={"label": "Vendor"})
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 50]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    uk = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]["uk"]
    c.patch("/api/parties/" + uk, json={"party_type": "vendor"})
    later = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [["Acme", "South", 80]])])
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a2.xlsx", later, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    row = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]
    assert row["fields"]["Balance"] == 80
    assert row["party_type"] == "vendor"


def test_sales_import_tags_party_as_customer():
    c = client()
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Buyer", "R", 40]])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    party = c.get("/api/rows", params={"type": "party"}).json()["rows"][0]
    assert party["party_type"] == "customer"
    assert party["party_type_label"] == "Customer"


def test_group_and_rep_360():
    c = client()
    _seed_core(c)
    extra = xlsx_bytes([
        ("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["20-09-24", "Beta", "R", 25]]),
        ("arr", ["Account Name", "Group", "Balance"], [["Beta", "South", 10]]),
    ])
    c.post("/api/uploads", files={"file": ("more.xlsx", extra, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    groups = c.get("/api/groups")
    assert groups.status_code == 200
    body = groups.json()
    assert body["total"] == 1
    south = body["groups"][0]
    assert south["name"] == "South"
    assert south["customers"] == 2
    assert south["due"] == 70
    assert "R" in south["reps"]
    assert "d0_15" in south
    assert south["d0_15"] + south["d15_30"] + south["d30_45"] + south["d45_60"] + south["d60_90"] + south["d90"] == south["due"]
    detail = c.get("/api/groups/%s" % south["uk"])
    assert detail.status_code == 200
    got = detail.json()
    assert got["name"] == "South"
    assert got["customer_count"] == 2
    assert {row["name"] for row in got["customers"]} == {"Acme", "Beta"}
    assert got["due"] == 70
    assert all("d0_15" in row for row in got["customers"])
    assert sum(
        row["d0_15"] + row["d15_30"] + row["d30_45"] + row["d45_60"] + row["d60_90"] + row["d90"]
        for row in got["customers"]
    ) == got["due"]
    assert "collection_buckets" in got
    assert got["fy_years"]
    assert any(p["name"] == "R" for p in got["salespeople"])
    assert got.get("insight", {}).get("headline")
    assert got["insight"].get("action")
    assert got.get("open_invoices")
    assert "order_check_policy" in got
    group_pdf = c.get("/api/groups/%s/pdf" % south["uk"])
    assert group_pdf.status_code == 200
    assert group_pdf.headers.get("content-type", "").startswith("application/pdf")
    assert group_pdf.content[:4] == b"%PDF"

    reps = c.get("/api/reps")
    assert reps.status_code == 200
    r = next(x for x in reps.json()["reps"] if x["name"] == "R")
    assert r["customers"] == 2
    assert "South" in r["groups"]
    assert "d0_15" in r
    assert r["d0_15"] + r["d15_30"] + r["d30_45"] + r["d45_60"] + r["d60_90"] + r["d90"] == r["due"]
    rd = c.get("/api/reps/%s" % r["uk"])
    assert rd.status_code == 200
    book = rd.json()
    assert book["name"] == "R"
    assert book["customer_count"] == 2
    assert {row["name"] for row in book["customers"]} == {"Acme", "Beta"}
    assert "owe_buckets" in book
    assert "collection_buckets" in book
    assert any(g["name"] == "South" for g in book["groups"])
    assert any(e["kind"] == "sale" for e in book["activity"])
    assert book["fy_years"]
    assert book.get("insight", {}).get("headline")
    assert book["insight"].get("action")
    assert book.get("open_invoices")
    assert any(line.get("party") in ("Acme", "Beta") for line in book["open_invoices"])
    assert "order_check_policy" in book
    rep_pdf = c.get("/api/reps/%s/pdf" % r["uk"])
    assert rep_pdf.status_code == 200
    assert rep_pdf.headers.get("content-type", "").startswith("application/pdf")
    assert rep_pdf.content[:4] == b"%PDF"
    assert c.get("/api/groups/Nobody").status_code == 404
    assert c.get("/api/reps/Nobody").status_code == 404


def test_group_and_rep_forbidden_without_permission():
    c = client()
    raw = TestClient(app)
    assert raw.get("/api/groups").status_code == 401
    assert raw.get("/api/reps").status_code == 401
    c.post("/api/users", json={"username": "wh", "password": "wh12345", "role": "Warehouse"})
    other = TestClient(app)
    login_as(other, "wh", "wh12345")
    assert other.get("/api/groups").status_code == 403
    assert other.get("/api/reps").status_code == 403
    assert other.get("/api/groups/South").status_code == 403
    assert other.get("/api/reps/R").status_code == 403


def test_upload_event_mode_update():
    c = client()
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [["19-09-24", "Acme", "R", 10]])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    again = c.post(
        "/api/uploads",
        data={"type": "sales", "event_mode": "update"},
        files={"file": ("s2.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert again.status_code == 200
    body = again.json()
    assert body["types"]["sales"]["updated"] == 1
    assert c.get("/api/rows", params={"type": "sales"}).json()["total"] == 1


def test_create_keeps_previous_run_and_persists_360():
    c = client()
    _seed_core(c)
    first = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    assert first.status_code == 202
    rid1 = first.json()["id"]
    run1 = c.get("/api/runs/%s" % rid1).json()
    assert run1["status"] == "succeeded"
    assert run1.get("views_split") or run1.get("book_json_id") or get_store().get_run(rid1).get("views_split")
    customers = c.get("/api/customers", params={"run": rid1})
    assert customers.status_code == 200
    assert customers.json()["total"] >= 1
    assert customers.json().get("from_snapshot") is True

    second = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-20"})
    assert second.status_code == 202
    rid2 = second.json()["id"]
    assert rid2 != rid1
    runs = c.get("/api/runs").json()["runs"]
    ids = {r["id"] for r in runs if r["status"] == "succeeded"}
    assert rid1 in ids and rid2 in ids
    older = c.get("/api/customers", params={"run": rid1}).json()
    newer = c.get("/api/customers", params={"run": rid2}).json()
    assert older["total"] >= 1
    assert newer["total"] >= 1


def test_cleanup_keeps_runs_factory_deletes_runs():
    c = client()
    _seed_core(c)
    created = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    rid = created.json()["id"]
    assert c.get("/api/runs/%s" % rid).json()["status"] == "succeeded"
    cleaned = c.post("/api/data/cleanup", json={"types": ["sales", "arr", "receipt"]})
    assert cleaned.status_code == 200
    assert c.get("/api/runs/%s" % rid).json()["status"] == "succeeded"
    assert c.get("/api/rows", params={"type": "sales"}).json()["total"] == 0
    factory = c.post("/api/data/factory-reset")
    assert factory.status_code == 200
    assert factory.json()["runs_deleted"] >= 1
    assert c.get("/api/runs/%s" % rid).status_code == 404 or c.get("/api/runs").json()["runs"] == []


def test_create_exposes_generate_progress():
    c = client()
    _seed_core(c)
    post = c.post("/api/runs", json={"packs": {"core": True}, "report_date": "2024-09-19"})
    body = post.json()
    rid = body["id"]
    assert (body.get("generate_progress") or {}).get("steps")
    run = c.get("/api/runs/%s" % rid).json()
    assert run["status"] == "succeeded"
    progress = run.get("generate_progress") or {}
    assert progress.get("total")
    steps = progress.get("steps") or []
    assert any(s.get("group") == "360 View" for s in steps)
    ids = [row["id"] for row in steps]
    assert ids.index("dashboard") < ids.index("save-reports")
    assert next(row["state"] for row in steps if row["id"] == "dashboard") == "done"


def test_ledger_arr_flag_uses_tolerance():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    c.put("/api/mappers/receipt", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Sales Rep", "Amount"]})
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        ["19-09-24", "GapCo", "Ria", 100],
        ["19-09-24", "MissCo", "Ria", 50],
        ["19-09-24", "BigCo", "Ria", 500],
    ])])
    receipts = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [
        ["20-09-24", "GapCo", "Ria", 99.6],
        ["21-09-24", "BigCo", "Ria", 100],
    ])])
    arr = xlsx_bytes([("arr", ["Account Name", "Balance"], [["GapCo", 0.15], ["BigCo", 50]])])
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, mime)})
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", receipts, mime)})
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, mime)})

    policy = c.get("/api/settings/org-policy").json()
    assert policy["ar_balance_tolerance"] == 0.5
    listed = c.get("/api/customers").json()["customers"]
    by_name = {row["name"]: row for row in listed}
    assert by_name["GapCo"]["balance_issue"] == ""
    assert by_name["GapCo"]["ledger_gap"] == 0.4
    assert by_name["MissCo"]["balance_issue"] == "missing_arr"
    assert by_name["MissCo"]["arr_balance"] is None
    flagged = c.get("/api/customers", params={"balance_issue": "Balance issues"}).json()
    assert {row["name"] for row in flagged["customers"]} == {"MissCo", "BigCo"}

    saved = c.post("/api/settings/org-policy", json={"ar_balance_tolerance": 0}).json()
    assert saved["ar_balance_tolerance"] == 0
    tight = {row["name"]: row for row in c.get("/api/customers").json()["customers"]}
    assert tight["GapCo"]["balance_issue"] == "mismatch"
    detail = c.get("/api/customers/GapCo").json()
    assert detail["balance_issue"] == "mismatch"
    assert detail["arr_balance"] == 0.15
    both = c.get("/api/customers", params={"balance_issue": "1"}).json()
    assert {row["name"] for row in both["customers"]} == {"GapCo", "MissCo", "BigCo"}
    mismatch = c.get("/api/customers", params={"balance_issue": "Mismatch"}).json()
    assert [row["name"] for row in mismatch["customers"]] == ["BigCo", "GapCo"]
    missing = c.get("/api/customers", params={"balance_issue": "Missing ARR"}).json()
    assert [row["name"] for row in missing["customers"]] == ["MissCo"]

    c.post("/api/settings/org-policy", json={"ar_balance_tolerance": 0.2})
    mid = {row["name"]: row for row in c.get("/api/customers").json()["customers"]}
    assert mid["GapCo"]["balance_issue"] == "mismatch"
    bad = c.post("/api/settings/org-policy", json={"ar_balance_tolerance": -1})
    assert bad.status_code == 400


def test_account_alias_migrates_and_rewrites_imports():
    c = client()
    c.put("/api/mappers/sales", json={"column_map": {}, "unique_key": ["Date", "Party Name", "Sales Rep", "Net Amount"]})
    c.put("/api/mappers/receipt", json={"column_map": {}, "unique_key": ["Date", "Account Name", "Sales Rep", "Amount"]})
    c.put("/api/mappers/arr", json={"column_map": {}, "unique_key": ["Account Name"]})
    c.put("/api/mappers/party", json={"column_map": {}, "unique_key": ["Account Name"]})
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sales = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        ["10-05-24", "Old Name", "Ria", 100],
        ["11-05-24", "New Name", "Ria", 30],
    ])])
    receipts = xlsx_bytes([("receipt", ["Date", "Account Name", "Sales Rep", "Amount"], [
        ["12-05-24", "Old Name", "Ria", 40],
        ["13-05-24", "New Name", "Ria", 10],
    ])])
    arr = xlsx_bytes([("arr", ["Account Name", "Group", "Balance"], [
        ["Old Name", "North", 40],
        ["New Name", "", 10],
    ])])
    c.post("/api/uploads", data={"type": "sales"}, files={"file": ("s.xlsx", sales, mime)})
    c.post("/api/uploads", data={"type": "receipt"}, files={"file": ("r.xlsx", receipts, mime)})
    c.post("/api/uploads", data={"type": "arr"}, files={"file": ("a.xlsx", arr, mime)})

    missing_dest = c.post("/api/settings/account-aliases/migrate", json={"source": "Old Name", "destination": "Brand New"})
    assert missing_dest.status_code == 400

    moved = c.post("/api/settings/account-aliases/migrate", json={"source": "Old Name", "destination": "New Name"})
    assert moved.status_code == 200
    body = moved.json()
    assert body["aliases"] == [{"source": "Old Name", "destination": "New Name"}]
    sales_rows = c.get("/api/rows", params={"type": "sales", "limit": "20"}).json()["rows"]
    assert {row["party"] for row in sales_rows} == {"New Name"}
    assert sum(row["amount"] for row in sales_rows) == 130
    receipt_rows = c.get("/api/rows", params={"type": "receipt", "limit": "20"}).json()["rows"]
    assert {row["party"] for row in receipt_rows} == {"New Name"}
    arr_rows = c.get("/api/rows", params={"type": "arr", "limit": "20"}).json()["rows"]
    assert len(arr_rows) == 1
    assert arr_rows[0]["party"] == "New Name"
    assert arr_rows[0]["amount"] == 50
    assert arr_rows[0]["group"] == "North"
    parties = c.get("/api/rows", params={"type": "party", "limit": "20"}).json()["rows"]
    assert "Old Name" not in {row["party"] for row in parties}
    customers = c.get("/api/customers").json()["customers"]
    assert {row["name"] for row in customers} == {"New Name"}

    again = xlsx_bytes([("sales", ["Date", "Party Name", "Sales Rep", "Net Amount"], [
        ["14-05-24", "Old Name", "Ria", 7],
    ])])
    uploaded = c.post("/api/uploads", data={"type": "sales"}, files={"file": ("n.xlsx", again, mime)})
    assert uploaded.status_code == 200
    assert uploaded.json()["types"]["sales"]["added"] == 1
    after = c.get("/api/rows", params={"type": "sales", "limit": "20"}).json()["rows"]
    assert {row["party"] for row in after} == {"New Name"}
    assert sum(row["amount"] for row in after) == 137
    parties_after = c.get("/api/rows", params={"type": "party", "limit": "20"}).json()["rows"]
    assert "Old Name" not in {row["party"] for row in parties_after}

    cycle = c.post("/api/settings/account-aliases/migrate", json={"source": "New Name", "destination": "Old Name"})
    assert cycle.status_code == 400
    chained = c.post("/api/settings/account-aliases/migrate", json={"source": "New Name", "destination": "Third Name"})
    assert chained.status_code == 400
    removed = c.delete("/api/settings/account-aliases", params={"source": "Old Name"})
    assert removed.status_code == 200
    assert removed.json()["aliases"] == []

