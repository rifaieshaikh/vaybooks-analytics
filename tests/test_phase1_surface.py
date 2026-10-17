"""Phase 1: frozen run basis, purchase cost history, and bridge surfaces."""

from io import BytesIO

from openpyxl import load_workbook

from server.customers import get_customer
from server.item_dims import get_item_dim
from server.items360 import get_item
from server.jobs import _ledger_store, enqueue_run, run_job
from server.ledger_surface import (
    bridge_reports,
    current_assessment,
    freeze_opening_basis,
    quality_rows,
    source_version,
)
from server.pdf_export import report_pdf_bytes
from server.persist import persist_type
from server.run_views import customer_profile
from server.store import get_store
from vay.export_excel import write_workbook


STOCK_MAP = {"column_map": {}, "unique_key": ["Item Name"]}


def _row(type_name, uk, fields, upload="u1"):
    return {
        "type": type_name,
        "uk": uk,
        "source_upload_id": upload,
        "fields": fields,
    }


def test_report_run_freezes_the_rows_it_will_read():
    store = get_store()
    store.upsert_row(_row("customer_opening", "open", {
        "Date": "2024-04-01", "Account Name": "Acme", "Balance": "0",
    }))
    store.upsert_row(_row("sales", "s1", {
        "Date": "2024-05-01", "Party Name": "Acme", "Sales Rep": "Ravi",
        "Net Amount": "100", "Invoice No": "INV-1",
    }))
    work = _ledger_store(store, [])
    assessment, basis = freeze_opening_basis(work, "2026-03-31")
    assert assessment["customers"][0]["customer_id"] == "acme"
    assert assessment["customers"][0]["bridge"]["invoices"] == "100.00"
    assert basis["policy"] == "invoice_date_fifo"
    assert basis["counts"]["sales"] == 1

    store.upsert_row(_row("sales", "s2", {
        "Date": "2024-06-01", "Party Name": "Acme", "Sales Rep": "Ravi",
        "Net Amount": "50", "Invoice No": "INV-2",
    }, upload="u2"))
    work._opening_assessment = None
    work._opening_basis = None
    frozen_again, frozen_basis = current_assessment(work, "2026-03-31")
    live, live_basis = current_assessment(store, "2026-03-31")
    assert frozen_basis["digest"] == basis["digest"]
    assert frozen_again["customers"][0]["bridge"]["invoices"] == "100.00"
    assert live["customers"][0]["bridge"]["invoices"] == "150.00"
    assert live_basis["digest"] != basis["digest"]
    assert "customer_opening" in source_version(work)["counts"]


def test_received_purchases_feed_the_bridge_without_changing_valuation():
    store = get_store()
    persist_type(
        store, "stock",
        ["Item Name", "Qty", "P.Price", "Supplier"],
        [["Pen", 10, 5, "Mill"]],
        STOCK_MAP, "stock-1",
    )
    store.upsert_row(_row("stock_opening", "open", {
        "Date": "2024-04-01", "Item Name": "Pen", "Qty": "6", "Unit": "pcs", "Value": "30",
    }))
    store.upsert_row(_row("purchases", "plan", {
        "Date": "2024-05-01", "Item Name": "Pen", "Qty": "2", "Unit": "pcs",
        "Invoice No": "PI-0", "Supplier": "Mill", "Rate": "8", "Receipt Status": "",
    }))
    store.upsert_row(_row("purchases", "pi", {
        "Date": "2024-05-02", "Item Name": "Pen", "Qty": "4", "Unit": "pcs",
        "Invoice No": "PI-1", "Movement Key": "M1", "Supplier": "Mill",
        "Rate": "7", "Amount": "28", "Receipt Status": "invoiced",
    }))
    store.upsert_row(_row("purchases", "gr", {
        "Date": "2024-05-03", "Item Name": "Pen", "Qty": "4", "Unit": "pcs",
        "Invoice No": "PI-1", "Movement Key": "M1", "Supplier": "Mill",
        "Rate": "9", "Amount": "36", "Receipt Status": "received",
    }))
    store.upsert_row(_row("purchases", "blank-rate", {
        "Date": "2024-05-04", "Item Name": "Pen", "Qty": "1", "Unit": "pcs",
        "Invoice No": "PI-2", "Supplier": "Mill", "Receipt Status": "received",
    }))

    pen = get_item(store, "Pen")
    history = pen["stock_bridge"]["cost_history"]
    assert [line["movement_key"] for line in history] == ["M1", "blank-rate"]
    assert history[0]["rate"] == "9.00"
    assert history[0]["amount"] == "36.00"
    assert history[1]["rate"] is None
    assert history[1]["amount"] is None
    assert pen["stock_bridge"]["valuation_method_unchanged"] is True
    bridges = pen["stock_bridge"]["bridges"]
    opened = next(row["bridge"] for row in bridges if row["bridge"].get("opening_qty") not in (None, ""))
    assert opened["opening_qty"] == "6.0000"
    assert opened["calculated_closing"] == "11.0000"
    assert store.find_row("stock", "Pen")["fields"]["P.Price"] == 5
    assert store.rows_of_type("item_cost") == []

    supplier = get_item_dim(store, "supplier", "Mill")
    purchases = supplier["purchases"]
    assert purchases["receipt_count"] == 2
    assert purchases["valuation_method_unchanged"] is True
    assert purchases["cost_history"][0]["rate"] == "9.00"
    assert store.find_row("stock", "Pen")["fields"]["P.Price"] == 5


def test_profiles_quality_and_exports_reuse_the_bridge():
    store = get_store()
    store.upsert_row(_row("customer", "acme", {"Account Name": "Acme", "Group": "South"}))
    store.upsert_row(_row("customer_opening", "open", {
        "Date": "2024-04-01", "Account Name": "Acme", "Balance": "80", "Invoice No": "OLD",
    }))
    store.upsert_row(_row("sales", "s1", {
        "Date": "2024-05-01", "Party Name": "Acme", "Sales Rep": "Ravi",
        "Net Amount": "20", "Invoice No": "INV-1",
    }))
    store.upsert_row(_row("arr", "acme-arr", {"Account Name": "Acme", "Balance": "100"}))
    store.upsert_row(_row("stock", "widget", {"Item Name": "Widget", "Qty": "3", "P.Price": "4"}))

    detail = get_customer(store, "Acme")
    bridge = detail["ledger_bridge"]
    assert bridge["opening_status"] == "imported"
    assert bridge["evidence"] == "source_confirmed"
    assert bridge["opening_amount"] == "80.00"
    assert "exact_balance" not in bridge["withheld_outputs"]
    profile = customer_profile(detail)
    assert profile["ledger_bridge"]["evidence"] == "source_confirmed"
    assert profile["sections"]["ledger"] is True

    assessment, _basis = current_assessment(store, detail["as_of"])
    rows = quality_rows(assessment)
    assert any("Widget" in row["message"] and "no imported stock opening" in row["message"] for row in rows)
    reports = bridge_reports(assessment)
    customer_sheet = reports[0]
    assert customer_sheet["rows"][0][0] == "acme"
    assert customer_sheet["rows"][0][2] == "source_confirmed"
    item_sheet = reports[1]
    assert item_sheet["rows"][0][0] == "Widget"
    assert "exact_balance" not in customer_sheet["headers"]
    assert "Withheld" in customer_sheet["headers"]

    book = load_workbook(BytesIO(write_workbook(reports)))
    assert "opening_customer_bridges" in book.sheetnames
    assert book["opening_customer_bridges"]["A2"].value == "acme"
    pdf = report_pdf_bytes(customer_sheet)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1000


def test_job_stores_the_frozen_basis_before_later_imports():
    store = get_store()
    store.upsert_row(_row("sales", "s1", {
        "Date": "2024-05-01", "Party Name": "Acme", "Sales Rep": "Ravi",
        "Net Amount": "25", "Invoice No": "INV-1",
    }))
    run_id, err = enqueue_run(store, {"core": True}, "2026-03-31", reports=["phase2_quality"])
    assert err == ""
    run_job(store, run_id)
    stored = store.get_run(run_id)
    assert stored["status"] == "succeeded"
    digest = stored["opening_basis"]["digest"]
    assert stored["opening_basis"]["policy"] == "invoice_date_fifo"
    assert stored["manifest"]["opening_basis"]["digest"] == digest
    assert stored["opening_assessment"]["customers"][0]["opening"]["status"] == "unresolved"
    titles = []
    blob = store.get_blob(stored["xlsx_id"])
    book = load_workbook(BytesIO(blob["data"]))
    titles.extend(book.sheetnames)
    assert "opening_customer_bridges" in titles

    store.upsert_row(_row("sales", "s2", {
        "Date": "2024-07-01", "Party Name": "Acme", "Sales Rep": "Ravi",
        "Net Amount": "10", "Invoice No": "INV-2",
    }, upload="later"))
    kept = store.get_run(run_id)
    assert kept["opening_basis"]["digest"] == digest
    assert source_version(store)["digest"] != digest
