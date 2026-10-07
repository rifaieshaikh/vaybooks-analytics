"""Excel statement, reorder proposal, and management summary match the screens."""

import os
from io import BytesIO

os.environ.setdefault("VAY_STORE", "memory")
os.environ.setdefault("VAY_SYNC_JOBS", "1")

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from server.actions import REVIEW_OPENS_UK, SETTING_TYPE
from server.customers import stated_money
from server.main import app
from server.org_policy import get_org_policy
from server.store import get_store, reset_store_for_tests


def client():
    reset_store_for_tests()
    c = TestClient(app)
    logged = c.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert logged.status_code == 200
    c.post("/api/settings/settlement", json={"mode": "oldest", "setup_complete": True})
    opened = c.post("/api/demo")
    assert opened.status_code == 200, opened.text
    return c


def sheet_rows(blob):
    book = load_workbook(BytesIO(blob))
    rows = []
    for ws in book.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append(tuple("" if value is None else value for value in row))
    return rows


def test_statement_excel_matches_the_customer_screen():
    c = client()
    listed = c.get("/api/customers").json()["customers"]
    harbour = next(row for row in listed if row["name"] == "Harbour Traders")
    detail = c.get("/api/customers/" + harbour["uk"]).json()
    policy = get_org_policy(get_store())
    detail["currency_code"] = policy.get("currency_code") or ""
    detail["currency_symbol"] = policy.get("currency_symbol") or ""
    downloaded = c.get("/api/customers/" + harbour["uk"] + "/xlsx")
    assert downloaded.status_code == 200, downloaded.text
    assert "spreadsheetml" in downloaded.headers["content-type"]
    rows = sheet_rows(downloaded.content)
    flat = {row[0]: row[1] for row in rows if len(row) >= 2 and row[0] in ("Company", "Customer", "Amount due", "Advance")}
    assert flat["Company"] == "Vay Demo Traders"
    assert flat["Customer"] == "Harbour Traders"
    due_label = "Advance" if detail.get("credit") else "Amount due"
    due_value = abs(detail.get("due") or 0) if detail.get("credit") else detail.get("due")
    assert flat[due_label] == stated_money(detail, due_value)
    aging = {row[0]: row[1] for row in rows if row and row[0] in ("0-15", "15-30", "30-45", "45-60", "60-90", "90+")}
    buckets = detail.get("owe_buckets") or {}
    assert aging["0-15"] == (buckets.get("d0_15") or 0)
    assert aging["90+"] == (buckets.get("d90") or 0)


def test_reorder_excel_matches_the_proposal():
    c = client()
    payload = c.get("/api/reorder").json()
    assert payload.get("status") == "eligible", payload
    lines = list(payload.get("order") or []) + list(payload.get("held") or [])
    oak = next(row for row in lines if row.get("name") == "Oak Board 18mm")
    downloaded = c.get("/api/reorder/xlsx")
    assert downloaded.status_code == 200, downloaded.text
    rows = sheet_rows(downloaded.content)
    match = next(row for row in rows if row and row[0] == "Oak Board 18mm")
    headers = next(row for row in rows if row and row[0] == "Item")
    suggested = match[headers.index("Suggested")]
    assert suggested == oak["suggested_qty"]
    formula = match[headers.index("Formula")]
    assert formula
    notes = " ".join(str(cell) for row in rows for cell in row)
    assert "not a purchase order" in notes


def test_review_excel_keeps_attribution_and_does_not_record_an_open():
    c = client()
    downloaded = c.get("/api/review/xlsx")
    assert downloaded.status_code == 200, downloaded.text
    assert get_store().find_row(SETTING_TYPE, REVIEW_OPENS_UK) is None
    rows = sheet_rows(downloaded.content)
    notes = " ".join(str(cell) for row in rows for cell in row)
    assert "not proof that the task caused the payment" in notes
    assert "does not show that the contact caused the sale" in notes
    review = c.get("/api/review").json()
    assert review["results"]["review_opens"] == 1
    again = c.get("/api/review/xlsx")
    assert again.status_code == 200
    assert c.get("/api/review").json()["results"]["review_opens"] == 1
